import numpy as np
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.ns import qn


def _gradient_letter(art_font, texture=False):
    font = art_font(100)
    fill = Image.new("L", (180, 160))
    outline = Image.new("L", fill.size)
    ImageDraw.Draw(fill).text((40, 10), "R", font=font, fill=255)
    ImageDraw.Draw(outline).text((40, 10), "R", font=font, fill=255, stroke_width=4, stroke_fill=255)
    mask, stroke = np.asarray(fill) > 127, np.asarray(outline) > 127
    pixels = np.full((160, 180, 3), (180, 220, 240), np.uint8)
    pixels[stroke] = (40, 30, 20)
    yy, xx = np.indices(mask.shape)
    amount = (yy-30)/70
    if texture:
        amount = ((xx//8+yy//8)%2).astype(float)
    colors = np.array([255, 245, 170])+amount[..., None]*np.array([-25, -95, -70])
    pixels[mask] = np.clip(colors[mask], 0, 255)
    item = {"text": "R", "font": font.getname()[0], "box": [0, 0, 180, 160],
            "words": [{"text": "R", "box": [0, 0, 1, 1]}]}
    return pixels, item


def test_linear_gradient_is_preserved_as_native_text_style(art_font):
    from scripts.art_text import estimate_art_text_runs
    pixels, item = _gradient_letter(art_font)
    runs = estimate_art_text_runs(pixels, item, reference_width=960)
    assert runs is not None
    gradient = runs[0].get("gradient")
    assert gradient is not None
    assert gradient["angle"] == pytest.approx(90, abs=2)
    colors = [[int(value[i:i+2], 16) for i in (1, 3, 5)] for value in gradient["colors"]]
    assert colors[0][1] - colors[1][1] > 80


def test_complex_fill_texture_is_not_accepted_as_a_linear_gradient(art_font):
    from scripts.art_text import estimate_art_text_runs
    pixels, item = _gradient_letter(art_font, texture=True)
    assert estimate_art_text_runs(pixels, item, reference_width=960) is None


def test_gradient_text_uses_drawingml_stops_and_remains_editable(tmp_path, art_font):
    from scripts import ppt_assemble as assembler
    deck = Presentation()
    item = {"text": "图", "box": [20, 20, 140, 140], "runs": [{
        "text": "图", "box": [0, 0, 1, 1], "box_kind": "ink", "font": art_font(100).getname()[0],
        "font_size": 100, "rotation": -12, "color": "#ffeaaa",
        "gradient": {"angle": 102, "colors": ["#fff0b0", "#e79965"]},
    }]}
    child = assembler._add_textbox(deck.slides.add_slide(deck.slide_layouts[6]), item,
                                   960, 540, assembler.compute_slide_transform(960, 540, "original")).shapes[0]
    rpr = child.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    assert rpr.find(qn("a:solidFill")) is None
    grad = rpr.find(qn("a:gradFill"))
    assert grad is not None
    assert list(rpr).index(grad) < list(rpr).index(rpr.find(qn("a:latin")))
    assert grad.find(qn("a:lin")).get("ang") == str(102*60000)
    assert [stop.find(qn("a:srgbClr")).get("val") for stop in grad.find(qn("a:gsLst"))] == ["fff0b0", "e79965"]
    path = tmp_path/"gradient.pptx"
    deck.save(path)
    reopened = Presentation(path)
    text = reopened.slides[0].shapes[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert text.text == "图"
    text.text = "园"
    reopened.save(path)
    assert Presentation(path).slides[0].shapes[0].shapes[0].text == "园"


@pytest.mark.parametrize("gradient", [{"angle": float("nan"), "colors": ["#ffffff", "#000000"]},
                                      {"angle": 90, "colors": ["#ffffff"]},
                                      {"angle": 90, "colors": ["broken", "#000000"]}])
def test_invalid_gradient_is_rejected_before_shape_creation(gradient):
    from scripts.text_runs import validate_text_runs
    with pytest.raises(ValueError, match="text runs"):
        validate_text_runs({"text": "A", "runs": [{"text": "A", "box": [0, 0, 1, 1], "gradient": gradient}]})


def test_bounded_gradient_does_not_extrapolate_beyond_observed_fill():
    from scripts.art_text import _linear_gradient
    yy, xx = np.indices((100, 30))
    region = (np.array([30, 100, 120]) + yy[..., None] * np.array([.5, 1, .8])).astype(np.uint8)
    interior = (yy >= 20) & (yy < 60)
    fill = np.ones((100, 30), dtype=bool)
    original = _linear_gradient(region, interior, fill, 0)
    assert original == _linear_gradient(region, interior, fill, 0, bound_samples=False)
    bounded = _linear_gradient(region, interior, fill, 0, bound_samples=True)
    assert bounded is not None
    colors = np.array([[int(value[i:i+2], 16) for i in (1, 3, 5)] for value in bounded['colors']])
    lo, hi = np.percentile(region[interior], (1, 99), axis=0)
    assert np.all(colors >= lo - 5.5)
    assert np.all(colors <= hi + 5.5)
    assert original['colors'] != bounded['colors']
