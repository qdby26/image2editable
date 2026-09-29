from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from scripts import font_embed


# Typefaces that this host's presentation app actually embeds: core system
# fonts (Arial, Calibri, ...) are skipped by the SaveAs embed flag.
_TYPEFACE_A = "Lucida Bright"
_TYPEFACE_B = "Noto Sans SC"


def _two_typeface_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 914400, 914400)
    run_a = box.text_frame.paragraphs[0].add_run()
    run_a.text = "First"
    run_a.font.name = _TYPEFACE_A
    paragraph = box.text_frame.add_paragraph()
    run_b = paragraph.add_run()
    run_b.text = "Second"
    run_b.font.name = _TYPEFACE_B
    presentation.save(path)


def _insert_embedded_font_entry(path: Path, *typefaces: str) -> None:
    staging = path.with_suffix(".staging.pptx")
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    presentation_xml = members["ppt/presentation.xml"].decode("utf-8")
    entry = "<p:embeddedFontLst>" + "".join(
        '<p:embeddedFont><p:font typeface="%s"/></p:embeddedFont>'
        % typeface
        for typeface in typefaces
    ) + "</p:embeddedFontLst>"
    assert "</p:presentation>" in presentation_xml
    members["ppt/presentation.xml"] = presentation_xml.replace(
        "</p:presentation>", entry + "</p:presentation>"
    ).encode("utf-8")
    with zipfile.ZipFile(staging, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    shutil.move(staging, path)


def test_collect_font_usage_reports_used_and_not_embedded(
    tmp_path: Path,
) -> None:
    deck = tmp_path / "deck.pptx"
    _two_typeface_deck(deck)
    _insert_embedded_font_entry(deck, _TYPEFACE_A)

    usage = font_embed.collect_font_usage(deck)

    assert usage["typefaces_used"] == [_TYPEFACE_A, _TYPEFACE_B]
    assert usage["embedded"] == [_TYPEFACE_A]
    assert usage["not_embedded"] == [_TYPEFACE_B]
    assert usage["size_bytes"] == deck.stat().st_size


def test_embed_fonts_requires_powerpoint(tmp_path: Path) -> None:
    if font_embed._powerpoint_available():
        pytest.skip("PowerPoint is installed here")
    with pytest.raises(RuntimeError, match="PowerPoint"):
        font_embed.embed_fonts(tmp_path / "a.pptx", tmp_path / "b.pptx")


def test_embed_fonts_embeds_used_typefaces(tmp_path: Path) -> None:
    if not font_embed._powerpoint_available():
        pytest.skip("win32com or PowerPoint dispatch unavailable")
    deck = tmp_path / "deck.pptx"
    embedded = tmp_path / "deck-embedded.pptx"
    _two_typeface_deck(deck)

    result = font_embed.embed_fonts(deck, embedded)

    assert embedded.is_file()
    assert result["dst_usage"]["not_embedded"] == []
    assert result["portable"] is True
    assert set(result["dst_usage"]["embedded"]) >= set(
        result["src_usage"]["typefaces_used"]
    )
    assert result["com_fonts"], "Presentation.Fonts listing expected"


def test_embed_pptx_in_place_replaces_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = tmp_path / "deck.pptx"
    _two_typeface_deck(deck)

    def fake_embed(src: Path, dst: Path) -> dict:
        shutil.copyfile(src, dst)
        return {
            "src_usage": font_embed.collect_font_usage(src),
            "dst_usage": font_embed.collect_font_usage(dst),
            "com_fonts": [],
            "app": {},
            "portable": True,
        }

    calls = []
    monkeypatch.setattr(
        font_embed, "embed_fonts",
        lambda s, d: calls.append((s, d)) or fake_embed(s, d),
    )
    payload = font_embed.embed_pptx_in_place(deck)

    assert calls and calls[0][0] == deck
    assert payload["embedded"] is True
    assert not (deck.parent / ".deck.embed-tmp.pptx").exists()
    report = deck.with_suffix(".embed-report.json")
    assert report.is_file()
    assert '"embedded": true' in report.read_text(encoding="utf-8")


def test_embed_pptx_in_place_disabled_by_env(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = tmp_path / "deck.pptx"
    _two_typeface_deck(deck)
    monkeypatch.setenv(font_embed.EMBED_FONTS_ENV, "0")
    monkeypatch.setattr(
        font_embed, "embed_fonts",
        lambda *_: pytest.fail("embed called while disabled"),
    )
    payload = font_embed.embed_pptx_in_place(deck)
    assert payload["skipped"] is True
    assert payload["embedded"] is False


def test_embed_pptx_in_place_failure_keeps_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = tmp_path / "deck.pptx"
    _two_typeface_deck(deck)
    digest = deck.read_bytes()

    def boom(*_):
        raise RuntimeError("COM unavailable")

    monkeypatch.setattr(font_embed, "embed_fonts", boom)
    payload = font_embed.embed_pptx_in_place(deck)

    assert payload["skipped"] is True
    assert "COM unavailable" in payload["reason"]
    assert deck.read_bytes() == digest
    assert deck.with_suffix(".embed-report.json").is_file()


def test_embed_pptx_in_place_already_embedded_short_circuits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = tmp_path / "deck.pptx"
    _two_typeface_deck(deck)
    _insert_embedded_font_entry(deck, _TYPEFACE_A, _TYPEFACE_B)
    monkeypatch.setattr(
        font_embed, "embed_fonts",
        lambda *_: pytest.fail("embed called on embedded deck"),
    )
    payload = font_embed.embed_pptx_in_place(deck)
    assert payload["embedded"] is True
    assert payload["already"] is True


def test_assemble_prepared_slide_embeds_delivery(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import image_to_ppt

    out = tmp_path / "out.pptx"
    slide_data = {
        "background_widescreen_path": "bg.png",
        "components": [],
        "text_items": [],
        "img_width": 4,
        "img_height": 4,
        "original_image_path": "src.png",
    }
    monkeypatch.setattr(
        image_to_ppt, "assemble_pptx",
        lambda **kwargs: (out.write_bytes(b"pptx"), str(out))[1],
    )
    calls = []
    monkeypatch.setattr(
        font_embed, "embed_pptx_in_place",
        lambda p, report_path=None: calls.append((p, report_path)) or {},
    )
    report = tmp_path / "named.embed-report.json"
    result = image_to_ppt._assemble_prepared_slide(
        slide_data, out, False, "16:9", embed_report_path=report,
    )
    assert result == str(out)
    assert calls == [(str(out), report)]
