"""Bundled font pool: matching candidates + embed substitutions."""
from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from scripts import font_embed, font_match

_POOL = Path(font_match.__file__).resolve().parents[1] / "fonts"


def test_installed_faces_includes_bundled_pool() -> None:
    names = {(f[0], f[1], f[2]) for f in font_match.installed_faces()}
    # Arimo VF registers upright + italic families; Tinos ships statics.
    assert ("Arimo", False, True) in names
    assert ("Tinos", False, False) in names
    assert ("Tinos", True, True) in names


def test_resolve_font_uses_pool_without_os_install() -> None:
    font = font_match.resolve_font("Arimo", italic=True)
    assert font is not None
    assert font.getname()[0] == "Arimo"


def test_substitute_map_targets_vendored_faces() -> None:
    families = {f[0].casefold() for f in font_match.installed_faces()}
    for core, substitute in font_embed.SUBSTITUTES.items():
        assert substitute.casefold() in families, (core, substitute)


def test_rewrite_typefaces_updates_runs(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Body"
    run.font.name = "Arial"
    presentation.save(deck)

    changed = font_embed._rewrite_typefaces(deck, {"Arial": "Arimo"})

    assert changed == {"Arial": "Arimo"}
    usage = font_embed.collect_font_usage(deck)
    assert usage["typefaces_used"] == ["Arimo"]


def test_embed_substitutes_core_font(tmp_path: Path) -> None:
    if not font_embed._powerpoint_available():
        pytest.skip("real PowerPoint dispatch unavailable")
    deck = tmp_path / "deck.pptx"
    embedded = tmp_path / "deck-embedded.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 914400, 914400)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Body text"
    run.font.name = "Arial"
    presentation.save(deck)

    result = font_embed.embed_fonts(deck, embedded)

    assert result["substitutions"] == {"Arial": "Arimo"}
    assert result["dst_usage"]["not_embedded"] == []
    assert result["portable"] is True
