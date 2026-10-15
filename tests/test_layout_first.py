import hashlib
import json
from pathlib import Path

import cv2
import numpy as np
import pytest
from PIL import Image
from pptx import Presentation

from image2editable import cli


_BANY06_FIXTURES = Path(__file__).parent / 'fixtures' / 'bany06'


def _bany06_fixture(kind, page, suffix=''):
    tag = f'{kind}-p{page:02d}{suffix}'
    source = _BANY06_FIXTURES / f'{tag}.png'
    cache = _BANY06_FIXTURES / f'{tag}.json'
    if not source.exists():
        source = _BANY06_FIXTURES / f'{kind}-p{page:02d}.png'
    return source, json.loads(cache.read_text(encoding='utf8'))['text_items']


@pytest.mark.parametrize('kind,page,phantom', [('run01', 3, 'O O'), ('run01', 3, 'O'), ('flatbg', 2, 'D目')])
def test_bany06_graphic_ocr_is_released_to_visual(kind, page, phantom, tmp_path):
    from image2editable.layout_first import prepare_layout
    source, items = _bany06_fixture(kind, page)
    slide, audit = prepare_layout(source, tmp_path/'assets', text_items=items,
                                  background_mode='source', label_reader=lambda *a, **k: ([], None))
    assert all(item['text'] not in ('O', 'O O', 'D目') for item in slide['text_items'])
    rejected = [item for item in items if item['text'] == phantom]
    rgb = np.asarray(Image.open(source).convert('RGB'))
    coverage = np.zeros(rgb.shape[:2], np.uint8)
    for element in slide['visual_elements']:
        c = element.get('component')
        if c:
            alpha = np.asarray(Image.open(c['path']).convert('RGBA'))[:, :, 3]
            coverage[c['y']:c['y']+c['h'], c['x']:c['x']+c['w']] |= alpha
    for item in rejected:
        x, y, w, h = map(round, item['box'])
        ink = rgb[y:y+h, x:x+w].min(axis=2) < 160
        assert (coverage[y:y+h, x:x+w][ink] > 0).mean() > .98
    if kind == 'run01':
        assert any(item['text'] == '3' for item in slide['text_items'])


def test_bany06_duplicate_narrow_prefix_is_not_native():
    from image2editable.layout_first import refine_layout_text
    source, items = _bany06_fixture('flatbg', 5)
    refined, _ = refine_layout_text(np.asarray(Image.open(source).convert('RGB')), items)
    assert any(item['text'].strip() == 'HPV' for item in refined)
    assert all(item['text'].strip() not in ('1', 'HPV 1') for item in refined)
    assert any('16/18' in item['text'] for item in refined)


@pytest.mark.parametrize('line', ['标准预防原则', '职业暴露处置流程', '手卫生规范', '隔离与防护用品'])
def test_bany06_one_ocr_line_is_not_split_by_ink_shading(line):
    from image2editable.layout_first import refine_layout_text
    source, items = _bany06_fixture('run01', 2)
    item = next(item for item in items if ''.join(item['text'].split()) == line)
    refined, _ = refine_layout_text(np.asarray(Image.open(source).convert('RGB')), [item])
    assert len(refined) == 1
    assert ''.join(refined[0]['text'].split()) == line


@pytest.mark.parametrize('kind,page,host,punctuation', [
    ('flatbg', 4, '平素体健', '，'), ('flatbg', 4, '不好意思就医', '，')])
def test_bany06_source_punctuation_follows_previous_text(kind, page, host, punctuation, tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    source, items = _bany06_fixture(kind, page)
    rgb = np.asarray(Image.open(source).convert('RGB'))
    selected = next(item for item in items if item['text'] == host)
    calls = []
    def reader(path, **kwargs):
        calls.append(path)
        pixels = np.asarray(Image.open(path).convert('RGB'))
        gray = cv2.cvtColor(pixels, cv2.COLOR_RGB2GRAY)
        mask = (gray < 160).astype(np.uint8)
        box = list(cv2.boundingRect(mask))
        return [{'text': host[-1]+punctuation, 'box': box, 'confidence': .99}], None
    refined, _, _ = _recover_text_gaps(rgb, [selected], _text_mask([selected], rgb.shape[:2]), tmp_path, reader)
    assert refined[0]['text'].endswith(punctuation)
    assert calls


def test_bany06_missing_middle_dot_matches_existing_source_dot(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask, _word_bounds
    source, items = _bany06_fixture('run01', 1)
    rgb = np.asarray(Image.open(source).convert('RGB'))
    original_bounds = _word_bounds(items[0])
    refined, _, _ = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path,
                                      lambda *a, **k: ([], None))
    json.dumps(refined)
    # the recovered dot is its own ink-positioned item at the source gap, not
    # text appended to the host run (inline flow lost the gap offset, B-any-07)
    assert refined[0]['text'] == '标准预防 · 手卫生'
    assert _word_bounds(refined[0]) == original_bounds
    host_right = refined[0]['box'][0] + refined[0]['box'][2]
    dot = next((item for item in refined if item['text'] == '·'
                and item['box'][0] > host_right), None)
    assert dot is not None
    assert dot.get('box_kind') == 'ink'
    assert refined[1]['text'] == '隔离防护'
    assert len(refined) == len(items)+1


def test_bany07_free_glyph_units_preserve_non_member_text_mask(tmp_path):
    from image2editable.layout_first import _free_glyph_units, _text_mask
    # the incoming mask is a union over items + stamps/labels/recovered entries;
    # rebuilding it from kept items alone dropped the subtitle band coverage and
    # left source glyphs under the native text (png2-p01 double image, v34+)
    rgb = np.full((600, 1200, 3), 255, np.uint8)
    items = [{'text': '常规文字', 'box': [50, 300, 200, 40]}]
    other = {'text': '《妇产科学》本科教学课件', 'box': [50, 400, 500, 60]}
    text = _text_mask(items + [other], rgb.shape[:2])
    visual_mask = np.zeros(rgb.shape[:2], np.uint8)
    kept, new_text, _, _ = _free_glyph_units(rgb, [], items, text.copy(), visual_mask)
    assert kept == items
    assert (new_text[400:460, 50:550] != 0).all()
    assert (new_text[300:340, 50:250] != 0).all()


def test_bany07_closing_punctuation_follows_host_prefix(tmp_path):
    from copy import deepcopy
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    source, items = _bany06_fixture('png2', 5)
    rgb = np.asarray(Image.open(source).convert('RGB'))
    for host_text, fragment, suffix in [('头号高危分子', '子”，', '”，'),
                                        ('头号高危分子', '危分子",', '”，'),
                                        ('(癌前病变)', '变)，', '，'),
                                        ('(癌前病变)', '前病变），', '，')]:
        host = deepcopy(next(item for item in items
                             if item['text'].rstrip().endswith(host_text)))
        tail = host['text'].rstrip()
        def reader(path, **kwargs):
            pixels = np.asarray(Image.open(path).convert('RGB'))
            if pixels.shape[1] > 500:
                # wide below-line windows read the next line's real text, not
                # the tail fragment — only the narrow tail crop sees it
                return [], None
            box = list(cv2.boundingRect((pixels.min(axis=2) < 160).astype(np.uint8)))
            return [{'text': fragment, 'box': box, 'confidence': .99}], None
        refined, _, _ = _recover_text_gaps(rgb, [host], _text_mask([host], rgb.shape[:2]),
                                          tmp_path, reader)
        if host_text == '头号高危分子':
            # the twin-mark+comma pattern resolves before OCR: the straight
            # closer '"' merges into the host flow (its marks live inside the
            # host box, so any standalone item would overlap G4) and '，'
            # becomes an ink item clipped past the host edge
            assert len(refined) == 2
            assert refined[0]['text'].endswith('"')
            punct = next(item for item in refined if item['text'] == '，')
            assert punct['box_kind'] == 'ink'
        else:
            # punctuation either merges into the host text or lands as an
            # ink-positioned item at the measured marks; it must never surface
            # as a duplicate text run reprinting the same glyphs
            if len(refined) == 1:
                assert refined[0]['text'] == tail + suffix
            else:
                assert len(refined) == 2
                punct = next(item for item in refined if item['text'] == suffix)
                assert punct['box_kind'] == 'ink'


def test_bany08_quote_tail_splits_straight_quote_and_ink_comma(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    source, items = _bany06_fixture('png2', 5)
    rgb = np.asarray(Image.open(source).convert('RGB'))
    host = next(item for item in items if item['text'].rstrip().endswith('头号高危分子'))
    # OCR never returns the bare tail; the unclosed '"' plus the
    # twin-mark+comma ink pattern splits: the source's straight closer '"'
    # joins the host flow (the marks sit inside the host box — any separate
    # item's frame would overlap it and trip G4), while '，' lands as an
    # ink item clipped just past the host's right edge.
    refined, _, _ = _recover_text_gaps(rgb, [host], _text_mask([host], rgb.shape[:2]),
                                      tmp_path, lambda *a, **k: ([], None))
    assert len(refined) == 2
    assert refined[0]['text'].endswith('"')
    punct = next(item for item in refined if item['text'] == '，')
    assert punct['box_kind'] == 'ink'
    x, y, w, h = refined[0]['box']
    px, pw = punct['box'][0], punct['box'][2]
    overlap = max(0, min(x + w, px + pw) - px)
    assert overlap <= .3 * min(h, punct['box'][3])
    assert 2600 <= px + pw / 2 <= 2650 and 750 <= punct['box'][1] <= 780


def test_bany08_straddling_mark_survives_text_mask_cleaning():
    from image2editable.layout_first import _preserve_straddling_marks
    # run03 p04: the ring '◯' at [3088,1156,20,19] dips its left arc under the
    # right edge of '接触血液体液时佩戴。' (frame ends 3095).  Box-shaped
    # cleaning repainted the arc and the surviving right half rendered as a
    # stray comma; mostly-outside ink is a neighbour mark, not a glyph.
    zone = np.full((60, 120, 3), 235, np.uint8)
    text = np.zeros((60, 120), np.uint8)
    text[:, 84:] = 255                       # sentence box ends mid-ring
    cv2.circle(zone, (80, 30), 10, (40, 40, 120), 3)   # ring ink x70..90
    cv2.rectangle(zone, (98, 22), (106, 38), (40, 40, 120), -1)  # '。' ink
    ys, xs = np.indices(zone.shape[:2])
    ring = (xs - 80) ** 2 + (ys - 30) ** 2 <= 196
    cleaning = (text != 0) & np.full((60, 120), 255, np.uint8).astype(bool)
    kept = _preserve_straddling_marks(cleaning.copy(), zone, text)
    ring_comp = ring & (zone.max(axis=2) < 160)
    glyph_comp = (zone.max(axis=2) < 160) & ~ring_comp
    # the ring survives whole; the glyph fully inside the box still cleans
    assert not np.any(kept & ring_comp)
    assert np.all(kept[glyph_comp])


def test_bany08_enclosed_weak_interior_fills_at_full_alpha(tmp_path):
    from image2editable.layout_first import prepare_layout
    # run03 p01 fill_lost [1760,480,320,280]: the white cross inside the
    # shield emblem scores dev~7 against the near-white page so the soft
    # alpha ramp left it at ~15% opacity and the render washed out.
    # Weak pixels fully enclosed by opaque ink are part of the object and
    # must come back at full alpha with their source colour.
    source, items = _bany06_fixture('run01', 1)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=items,
                              background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    pics = [e for e in slide['visual_elements']
            if e['route'] == 'raster_component' and e.get('component')]
    cx, cy = 1920, 620                          # inside the white cross
    target = [e for e in pics
              if e['component']['x'] <= cx < e['component']['x'] + e['component']['w']
              and e['component']['y'] <= cy < e['component']['y'] + e['component']['h']]
    assert target, [e['component'] for e in pics]
    rgba = np.asarray(Image.open(target[0]['component']['path']))
    local = rgba[cy - target[0]['component']['y'],
                 cx - target[0]['component']['x']]
    assert local[3] > 200, f'cross alpha {local[3]}'
    assert local[:3].min() > 240, f'cross rgb {local[:3]}'


def test_bany08_subtle_fill_block_emits_native_shape(tmp_path):
    from image2editable.layout_first import prepare_layout
    # Cluster C-1: a uniform block within 5..20 of the backdrop (icon
    # underlay, pale case card, white chip on a tinted page) sits under the
    # >20 residual floor and collapses into the flat background.  It must
    # emit as a native fill shape instead.
    rgb = np.full((900, 1500, 3), (247, 250, 255), np.uint8)
    cv2.rectangle(rgb, (300, 300), (900, 700), (235, 243, 253), -1)
    source = tmp_path / 'card.png'
    Image.fromarray(rgb).save(source)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=[],
                              label_reader=lambda *a, **k: ([], None))
    blocks = [e for e in slide['visual_elements']
              if e['route'] == 'native_shape' and e.get('shape', {}).get('fill_rgb')]
    hit = next((e for e in blocks
                if e['bbox'][0] <= 320 and e['bbox'][2] >= 880
                and e['bbox'][1] <= 320 and e['bbox'][3] >= 680
                and max(abs(a - b) for a, b in zip(
                        e['shape']['fill_rgb'], (235, 243, 253))) <= 8), None)
    assert hit is not None, [(e['bbox'], e['shape'].get('fill_rgb')) for e in blocks]


def test_bany08_card_edge_centres_on_border_strip():
    from image2editable.layout_first import _border_band_centre
    # run03 p02: the source border is an ~8px gradient strip peaking inside
    # the mask while the mask edge overshoots ~3px; centring the stroke on
    # the mask edge rendered the line a couple px low and the detector saw
    # broken segments.  The shape edge must track the measured band centre.
    cover = np.zeros(900, bool)
    cover[699:707] = True                            # strip straddling edge 706
    centre, thick = _border_band_centre(cover, 706, 10)
    assert abs(centre - 702.5) <= .5 and thick == 8
    inward = np.zeros(900, bool)
    inward[696:702] = True                           # strip ending before edge
    centre, thick = _border_band_centre(inward, 706, 10)
    assert abs(centre - 698.5) <= .5 and thick == 6
    outward = np.zeros(900, bool)
    outward[708:716] = True                          # strip just past the edge
    centre, thick = _border_band_centre(outward, 706, 10)
    assert abs(centre - 711.5) <= .5 and thick == 8
    assert _border_band_centre(np.zeros(900, bool), 300, 10) is None
    far = np.zeros(900, bool)
    far[200:210] = True                              # distant ink is not a band
    assert _border_band_centre(far, 706, 10) is None


def test_bany07_sparse_residual_over_text_is_dropped(tmp_path):
    from image2editable.layout_first import _emit_residual
    difference = np.full((2160, 3840), 100, np.uint8)
    text = np.zeros((2160, 3840), np.uint8)
    items = [{'text': '黄金处置时间：', 'box': [1700, 1650, 300, 90]}]
    cv2.rectangle(text, (1700, 1650), (2000, 1740), 255, -1)

    def _sparse(width, height, dots=6):
        rgba = np.zeros((height, width, 4), np.uint8)
        for i in range(dots):
            rgba[3 + i * 5 % height, 3 + i * 7 % width] = [80, 120, 200, 255]
        return rgba
    elements = []
    records = [
        {'rgba': _sparse(560, 100), 'x': 1296, 'y': 1646, 'w': 560, 'h': 100, 'owned': False},
        {'rgba': _sparse(200, 90), 'x': 3000, 'y': 1646, 'w': 200, 'h': 90, 'owned': False},
    ]
    dense = np.zeros((90, 200, 4), np.uint8)
    dense[10:80, 20:180] = [200, 60, 60, 255]
    records.append({'rgba': dense, 'x': 1760, 'y': 1660, 'w': 200, 'h': 90, 'owned': False})
    count = _emit_residual(records, elements, tmp_path, difference, (2160, 3840), text, items)
    boxes = [(e['component']['x'], e['component']['y']) for e in elements]
    assert (3000, 1646) in boxes
    assert (1760, 1660) in boxes
    assert count == 2
    assert all(not (e['component']['x'] == 1296 and e['component']['y'] == 1646)
               for e in elements)


def test_bany06_superscript_correction_preserves_native_numbered_labels():
    from image2editable.layout_first import refine_layout_text
    rgb = np.full((100, 300, 3), 255, np.uint8)
    items = [{'text': '1', 'box': [10, 20, 12, 40]},
             {'text': '16/18型', 'box': [40, 20, 160, 40]}]
    refined, _ = refine_layout_text(rgb, items)
    assert [item['text'] for item in refined] == ['1', '16/18型']


def test_bany06_middle_comma_uses_real_same_line_neighbors(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    source, items = _bany06_fixture('flatbg', 4)
    rgb = np.asarray(Image.open(source).convert('RGB'))
    selected = [next(item for item in items if item['text'] == label)
                for label in ('不好意思就医', '自购妇科洗液冲洗；症状反复2')]
    def reader(path, **kwargs):
        pixels = np.asarray(Image.open(path).convert('RGB'))
        box = list(cv2.boundingRect((pixels.min(axis=2) < 160).astype(np.uint8)))
        return [{'text': '医，', 'box': box, 'confidence': .99}], None
    refined, _, _ = _recover_text_gaps(rgb, selected, _text_mask(selected, rgb.shape[:2]), tmp_path, reader)
    assert refined[0]['text'] == '不好意思就医，'
    assert len(refined) == 2


@pytest.mark.parametrize('read_prefix', [False, True])
def test_unhosted_percent_plus_is_one_raster_unit(tmp_path, read_prefix):
    from image2editable.layout_first import prepare_layout

    rgb = np.full((600, 1200, 3), 255, np.uint8)
    cv2.putText(rgb, '70%', (80, 350), cv2.FONT_HERSHEY_SIMPLEX, 4, (20, 65, 180), 10)
    cv2.putText(rgb, '+', (350, 350), cv2.FONT_HERSHEY_SIMPLEX, 3, (20, 65, 180), 10)
    source = tmp_path/'percent.png'
    Image.fromarray(rgb).save(source)
    items = [{'text': '70', 'box': [80, 260, 160, 96], 'color': '#1441b4'}] if read_prefix else []
    def reader(path, **kwargs):
        if 'glyph-input' not in Path(path).name:
            return [], None
        return [{'text': '70%+', 'confidence': .99}], None
    slide, _ = prepare_layout(source, tmp_path/'assets', text_items=items,
                              label_reader=reader)
    parts = [e['component'] for e in slide['visual_elements'] if e.get('component')
             and e['component']['y'] < 360 and e['component']['y']+e['component']['h'] > 260]
    assert len(parts) == 1, parts
    assert not slide['text_items']
    c = parts[0]
    assert c['x'] < 100 and c['x']+c['w'] > 410


@pytest.mark.parametrize('bar_x', [435, 465])
def test_free_glyph_unit_does_not_absorb_neighboring_divider(tmp_path, bar_x):
    from image2editable.layout_first import prepare_layout

    rgb = np.full((600, 1200, 3), 255, np.uint8)
    cv2.putText(rgb, '70%', (80, 350), cv2.FONT_HERSHEY_SIMPLEX, 4, (20, 65, 180), 10)
    cv2.putText(rgb, '+', (350, 350), cv2.FONT_HERSHEY_SIMPLEX, 3, (20, 65, 180), 10)
    rgb[270:385, bar_x:bar_x+4] = (180, 180, 180)
    source = tmp_path/'percent-divider.png'
    Image.fromarray(rgb).save(source)
    def reader(path, **kwargs):
        return ([{'text': '70%+', 'confidence': .99}], None) if 'glyph-input' in Path(path).name else ([], None)
    slide, _ = prepare_layout(source, tmp_path/'assets', text_items=[], label_reader=reader)
    unit = next(e for e in slide['visual_elements'] if e['object_id'].startswith('unit_'))
    comp = unit['component']
    rgba = np.asarray(Image.open(comp['path']).convert('RGBA'))
    gray = (rgba[:, :, :3].max(axis=2) == rgba[:, :, :3].min(axis=2)) & (rgba[:, :, 0] == 180)
    assert not (gray & (rgba[:, :, 3] > 0)).any()
    coverage = np.zeros(rgb.shape[:2], np.uint8)
    for element in slide['visual_elements']:
        if element is unit or not element.get('component'):
            continue
        c = element['component']
        alpha = np.asarray(Image.open(c['path']).convert('RGBA'))[:, :, 3]
        coverage[c['y']:c['y']+c['h'], c['x']:c['x']+c['w']] |= alpha
    assert (coverage[270:385, bar_x:bar_x+4] > 0).all()


@pytest.mark.parametrize('contained', [False, True])
def test_confirmed_glyph_unit_rejects_foreign_residual_sliver(tmp_path, contained):
    from image2editable.layout_first import _emit_residual

    rgb = np.full((120, 350, 3), 255, np.uint8)
    cv2.putText(rgb, '70%+', (0, 90), cv2.FONT_HERSHEY_SIMPLEX, 2, (20, 65, 180), 5)
    alpha = (rgb.min(axis=2) < 255).astype(np.uint8)*255
    path = tmp_path/'unit_0001.png'
    Image.fromarray(np.dstack((rgb, alpha))).save(path)
    unit = {'object_id': 'unit_0001', 'route': 'raster_component', 'role': 'visual',
            'residual_merge': False, 'component': {'path': str(path), 'x': 80, 'y': 260, 'w': 350, 'h': 120}}
    before = path.read_bytes()
    bar = np.full((115, 4, 4), (180, 180, 180, 255), np.uint8)
    records = [{'x': 420 if contained else 440, 'y': 270, 'w': 4, 'h': 115, 'rgba': bar, 'owned': False}]
    elements = [unit]
    blank = np.zeros((1200, 2400), np.uint8)
    _emit_residual(records, elements, tmp_path, blank, blank.shape, blank, [])
    assert path.read_bytes() == before
    assert len(elements) == 2
    np.testing.assert_array_equal(np.asarray(Image.open(elements[1]['component']['path'])), bar)


def test_three_parallel_arrows_are_not_confirmed_as_text(tmp_path):
    from image2editable.layout_first import _free_glyph_units

    rgb = np.full((600, 1200, 3), 255, np.uint8)
    for x in (80, 160, 240):
        cv2.fillPoly(rgb, [np.array([[x, 200], [x+30, 200], [x+65, 240],
                                    [x+30, 280], [x, 280], [x+35, 240]])], (20, 65, 180))
    blank = np.zeros((600, 1200), np.uint8)
    _, _, units, _ = _free_glyph_units(rgb, [], [], blank, blank)
    assert units == []


def test_nested_cards_move_with_outer_card_in_both_modes(tmp_path):
    from image2editable.layout_first import _group_layout
    for mode in ('native', 'png'):
        elements = []
        for index, box in enumerate(([20, 20, 220, 160], [60, 60, 90, 100]), 1):
            x, y, right, bottom = box
            path = tmp_path / f'{mode}-{index}.png'
            Image.new('RGBA', (right-x, bottom-y), 'blue').save(path)
            elements.append({'object_id': f'card_{index}', 'role': 'card', 'route': 'raster_component',
                             'z_index': index, 'bbox': box,
                             'component': {'path': str(path), 'x': x, 'y': y, 'w': right-x, 'h': bottom-y}})
        items = [{'text': 'label', 'box': [70, 70, 10, 15]}]
        _, groups = _group_layout(elements, items, tmp_path, (400, 240), mode)
        assert len(groups) == 1
        assert 'text_0001' in next(iter(groups.values()))
        if mode == 'native':
            assert {'card_1', 'card_2'} <= set(next(iter(groups.values())))


def test_frame_groups_contained_text_and_border_lines_without_swallowing_crossing_visual(tmp_path):
    from image2editable.layout_first import _group_layout
    elements = [
        {'object_id': 'frame', 'role': 'frame', 'route': 'native_shape', 'bbox': [30, 30, 230, 130]},
        {'object_id': 'line', 'role': 'line', 'route': 'native_shape', 'bbox': [35, 32, 225, 35]},
        {'object_id': 'icon', 'role': 'visual', 'bbox': [50, 50, 80, 80]},
        {'object_id': 'crossing', 'role': 'visual', 'bbox': [0, 45, 300, 95]},
    ]
    items = [{'text': 'body', 'box': [100, 50, 100, 40]}]
    _, groups = _group_layout(elements, items, tmp_path, (400, 240), 'native')
    assert len(groups) == 1
    assert set(next(iter(groups.values()))) == {'frame', 'line', 'icon', 'text_0001'}


def test_adjacent_frame_and_footer_remain_independently_movable(tmp_path):
    from image2editable.layout_first import _group_layout
    elements = [{'object_id': 'frame', 'role': 'frame', 'bbox': [20, 20, 220, 120]},
                {'object_id': 'footer', 'role': 'card', 'bbox': [20, 123, 220, 150]}]
    _, groups = _group_layout(elements, [], tmp_path, (400, 240), 'native')
    assert len(groups) == 2


def test_stacked_gutter_cards_remain_independently_movable(tmp_path):
    from image2editable.layout_first import _group_layout
    elements = [{'object_id': 'card_a', 'role': 'card', 'bbox': [40, 20, 380, 120]},
                {'object_id': 'card_b', 'role': 'card', 'bbox': [40, 124, 380, 224]},
                {'object_id': 'badge', 'role': 'card', 'bbox': [50, 140, 90, 200]}]
    items = [{'text': 'A', 'box': [60, 40, 80, 20]}, {'text': 'B', 'box': [60, 150, 80, 20]}]
    _, groups = _group_layout(elements, items, tmp_path, (400, 260), 'native')
    assert len(groups) == 2
    members = list(groups.values())
    assert any({'card_a', 'text_0001'} <= set(ids) for ids in members)
    assert any({'card_b', 'badge', 'text_0002'} <= set(ids) for ids in members)


def test_frame_partitions_crossing_raster_without_losing_pixels(tmp_path):
    from image2editable.layout_first import _group_layout
    rgba = np.full((140, 300, 4), 255, np.uint8)
    rgba[40:90, 100:210, :3] = [20, 50, 90]
    path = tmp_path / 'crossing.png'
    Image.fromarray(rgba).save(path)
    elements = [{'object_id': 'frame', 'role': 'frame', 'bbox': [90, 30, 230, 110]},
                {'object_id': 'crossing', 'role': 'visual', 'route': 'raster_component', 'z_index': 1000,
                 'component': {'path': str(path), 'x': 0, 'y': 0, 'w': 300, 'h': 140}}]
    result, groups = _group_layout(elements, [], tmp_path, (400, 240), 'native')
    members = set(next(iter(groups.values())))
    assert 'crossing' not in members
    assert len(members) == 2
    reconstructed = Image.new('RGBA', (300, 140))
    alpha_sum = np.zeros((140, 300), dtype=int)
    for element in result:
        if element.get('role') != 'visual':
            continue
        c = element['component']
        with Image.open(c['path']) as patch:
            reconstructed.alpha_composite(patch, (c['x'], c['y']))
            alpha_sum[c['y']:c['y']+c['h'], c['x']:c['x']+c['w']] += np.asarray(patch)[:, :, 3]
    assert np.array_equal(np.asarray(reconstructed), rgba)
    assert np.all(alpha_sum == 255)


def test_enclosing_frame_preserves_editable_child_card_groups(tmp_path):
    from image2editable.layout_first import _group_layout
    elements = [{'object_id': 'frame', 'role': 'frame', 'bbox': [10, 10, 310, 210]},
                {'object_id': 'label_a', 'role': 'card', 'bbox': [30, 30, 90, 60]},
                {'object_id': 'label_b', 'role': 'card', 'bbox': [30, 100, 90, 130]}]
    items = [{'text': 'A', 'box': [40, 35, 20, 15]}, {'text': 'B', 'box': [40, 105, 20, 15]}]
    _, groups = _group_layout(elements, items, tmp_path, (400, 240), 'native')
    assert len(groups) == 3
    a = next(g for g, ids in groups.items() if 'label_a' in ids)
    b = next(g for g, ids in groups.items() if 'label_b' in ids)
    parent = next(g for g, ids in groups.items() if 'frame' in ids)
    assert 'text_0001' in groups[a] and 'text_0002' in groups[b]
    assert a in groups[parent] and b in groups[parent]
    assert list(groups).index(parent) > list(groups).index(a)


def test_assembler_retains_child_groups_inside_frame(tmp_path):
    from scripts.ppt_assemble import assemble_pptx_multi
    output = tmp_path / 'nested.pptx'
    slide = {'img_width': 400, 'img_height': 240, 'components': [],
             'background_rgb': [255, 255, 255],
             'text_items': [{'text': 'A', 'box': [40, 40, 30, 20], 'layout_object_id': 'text_a'},
                            {'text': 'B', 'box': [40, 100, 30, 20], 'layout_object_id': 'text_b'}],
             'object_groups': {'child_a': ['text_a'], 'child_b': ['text_b'], 'parent': ['child_a', 'child_b']}}
    assemble_pptx_multi([slide], output, add_reference=False)
    prs = Presentation(output)
    assert len(prs.slides[0].shapes) == 1
    parent = prs.slides[0].shapes[0]
    assert parent.name == 'image2editable:parent'
    assert [child.name for child in parent.shapes] == ['image2editable:child_a', 'image2editable:child_b']


def test_nonintersecting_visual_does_not_decode_image_for_partition(tmp_path):
    from image2editable.layout_first import _group_layout
    visual = {'object_id': 'icon', 'role': 'visual', 'route': 'raster_component', 'z_index': 1000,
              'component': {'path': str(tmp_path / 'not-opened.png'), 'x': 250, 'y': 150, 'w': 40, 'h': 40}}
    frame = {'object_id': 'frame', 'role': 'frame', 'bbox': [10, 10, 110, 110]}
    result, _ = _group_layout([frame, visual], [], tmp_path, (400, 240), 'native')
    assert visual in result


def test_badge_owns_raster_digit_before_enclosing_frame_partition(tmp_path):
    from image2editable.layout_first import _group_layout
    path = tmp_path / 'badge-digit.png'
    rgba = np.full((160, 300, 4), 255, np.uint8)
    rgba[45:65, 45:55, :3] = [20, 40, 200]
    Image.fromarray(rgba).save(path)
    elements = [{'object_id': 'frame', 'role': 'frame', 'bbox': [10, 10, 230, 110]},
                {'object_id': 'badge', 'role': 'card', 'bbox': [30, 30, 80, 80]},
                {'object_id': 'art', 'role': 'visual', 'route': 'raster_component', 'z_index': 1000,
                 'component': {'path': str(path), 'x': 0, 'y': 0, 'w': 300, 'h': 160}}]
    result, groups = _group_layout(elements, [], tmp_path, (400, 240), 'native')
    badge_group = next(ids for ids in groups.values() if 'badge' in ids)
    assert len(badge_group) == 2
    owned = next(e for e in result if e['object_id'] in badge_group and e['role'] == 'visual')
    assert owned['component']['x'] == 30 and owned['component']['y'] == 30
    reconstructed = Image.new('RGBA', (300, 160))
    for e in result:
        if e['role'] == 'visual':
            c = e['component']
            with Image.open(c['path']) as patch:
                reconstructed.alpha_composite(patch, (c['x'], c['y']))
    assert np.array_equal(np.asarray(reconstructed), rgba)


@pytest.mark.parametrize('outline', [(232, 236, 241), (227, 234, 247)])
def test_low_chroma_outline_card_is_detected_and_groups_body(tmp_path, outline):
    from PIL import ImageDraw
    from image2editable.layout_first import prepare_layout
    source = tmp_path / 'pale-outline.png'
    image = Image.new('RGB', (800, 360), 'white')
    ImageDraw.Draw(image).rounded_rectangle((50, 90, 750, 170), radius=12, outline=outline, width=2)
    image.save(source)
    slide, audit = prepare_layout(source, tmp_path / 'assets', text_items=[{'text': 'body', 'box': [150, 110, 400, 40]},
                                                                       {'text': 'title', 'box': [30, 60, 500, 40]}])
    assert any(e['role'] == 'frame' for e in slide['visual_elements'])
    assert any('text_0001' in ids for ids in audit['object_groups'].values())


def test_vertical_surface_gradient_uses_powerpoint_top_to_bottom_angle():
    from image2editable.layout_first import _surface
    colors = np.linspace([30, 110, 230], [30, 50, 180], 120).astype(np.uint8)
    rgb = np.broadcast_to(colors[:, None, :], (120, 200, 3)).copy()
    mask = np.full((120, 200), 255, np.uint8)
    _, payload = _surface(rgb, mask, np.zeros_like(mask))
    assert payload['gradient']['angle'] == 270
    assert np.max(np.abs(np.asarray(payload['gradient']['colors'][0]) - rgb[0, 100])) <= 2
    assert np.max(np.abs(np.asarray(payload['gradient']['colors'][1]) - rgb[-1, 100])) <= 2


def test_shared_top_rule_is_not_attached_to_one_of_several_frames(tmp_path):
    from image2editable.layout_first import _group_layout
    elements = [{'object_id': 'left', 'role': 'frame', 'bbox': [20, 40, 100, 150]},
                {'object_id': 'right', 'role': 'frame', 'bbox': [120, 40, 200, 150]},
                {'object_id': 'rule', 'role': 'line', 'bbox': [20, 37, 200, 40]}]
    _, groups = _group_layout(elements, [], tmp_path, (400, 240), 'native')
    assert len(groups) == 2
    assert all('rule' not in ids for ids in groups.values())


def test_bright_runs_sample_fill_inside_antialiased_ink(monkeypatch):
    from image2editable.layout_first import _bright_title_runs
    import scripts.font_match as font_match
    monkeypatch.setattr(font_match, 'match_glyph', lambda *args: {
        'font': 'Microsoft YaHei', 'bold': True, 'rotation': 0,
        'font_size': 64, 'fit_iou': .95})
    rgb = np.full((80, 100, 3), (30, 70, 130), np.uint8)
    rgb[20:60, 30:70] = (210, 220, 240)
    rgb[22:58, 32:68] = 255
    item = {'text': '“', 'box': [0, 0, 100, 80]}
    runs = _bright_title_runs(rgb, item, 3840)
    assert runs is not None and runs[0]['color'] == '#ffffff'
    rgb[22:40, 32:68] = (210, 240, 255)
    assert _bright_title_runs(rgb, item, 3840) is None
    rgb[:] = (30, 70, 130)
    rgb[20:60, 30:32] = 255
    assert _bright_title_runs(rgb, item, 3840) is None


def test_recovered_fragment_runs_use_complete_ink_components():
    from PIL import ImageDraw
    from scripts.font_match import resolve_font
    from image2editable.layout_first import _recovered_fragment_runs
    font = resolve_font('Microsoft YaHei', size=64)
    if font is None:
        pytest.skip('Requires Microsoft YaHei')
    image = Image.new('RGB', (320, 120), 'white')
    ImageDraw.Draw(image).text((20, 10), '(CT、', font=font, fill=(30, 30, 30))
    rgb = np.asarray(image)
    box = list(cv2.boundingRect((rgb.min(axis=2) < 140).astype(np.uint8)))
    item = {'text': '(CT、', 'box': box, 'font': 'Arial', 'font_size': 16}
    runs = _recovered_fragment_runs(rgb, item, 3840)
    assert ''.join(run['text'] for run in runs) == item['text']
    assert all(run['box_kind'] == 'ink' for run in runs)
    assert runs[-1]['font'] == 'Microsoft YaHei'
    assert all(run['font_size'] > 14 for run in runs)
    json.dumps(runs)
    assert _recovered_fragment_runs(np.full_like(rgb, 255), item, 3840) is None
    assert _recovered_fragment_runs(rgb, dict(item, text='(CT、X'), 3840) is None
    assert _recovered_fragment_runs(255-rgb, item, 3840) is None
    assert _recovered_fragment_runs(rgb, dict(item, text='吗？'), 3840) is None
    assert _recovered_fragment_runs(rgb, dict(item, text='(OC、'), 3840) is None
    assert _recovered_fragment_runs(rgb, dict(item, text='CT'), 3840) is None


def test_text_gap_recovery_maps_boxes_and_rejects_invalid_results(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': 'left', 'box': [60, 55, 150, 40]},
             {'text': 'right', 'box': [340, 55, 150, 40]}]
    calls = []
    def reader(path, **kwargs):
        calls.append(path)
        return [{'text': 'CT', 'box': [5, 5, 45, 30], 'confidence': .99}], None
    result, mask, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert len(calls) == 1 and audit['gap_texts_recovered'] == 1
    assert 210 < result[-1]['box'][0] < 300
    json.dumps(result)
    assert mask[round(result[-1]['box'][1])+10, round(result[-1]['box'][0])+10] == 255
    for found in ([{'text': 'CT', 'box': [-1, 0, 45, 30], 'confidence': .99}],
                  [{'text': 'CT', 'box': [5, 5, 45, 30], 'confidence': .7}]):
        unchanged, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, lambda *a, **k: (found, None))
        assert unchanged == items and audit['gap_texts_recovered'] == 0


@pytest.mark.parametrize('confirmed,confidence,fit', [('“点状”', .99, True), ('“点状”', .97, True), ('“其他”', .99, True), ('“点状”', .99, False)])
def test_gap_quote_retry_requires_agreement_confidence_and_glyph_fit(tmp_path, monkeypatch, confirmed, confidence, fit):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': 'left', 'box': [60, 55, 150, 40]}]
    calls = []
    def reader(path, **kwargs):
        calls.append(kwargs)
        if kwargs.get('recognition_only'):
            width, height = Image.open(path).size
            return [{'text': confirmed, 'box': [0, 0, width, height], 'confidence': confidence}], None
        return [{'text': '点状”', 'box': [5, 5, 45, 30], 'confidence': .96}], None
    monkeypatch.setattr('image2editable.layout_first._bright_title_runs', lambda *args: [{'text': '“', 'box': [.1, .1, .2, .5]}] if fit else None)
    result, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    accepted = confirmed == '“点状”' and confidence >= .98 and fit
    assert result[-1]['text'] == ('“点状”' if accepted else '点状”')
    assert len(calls) == 2 and audit['gap_recovery_calls'] == 2
    assert calls[-1]['recognition_only']


def test_quote_retry_preserves_original_when_new_runs_overlap_existing_text(tmp_path, monkeypatch):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': 'left', 'box': [60, 55, 150, 40]}]
    text = _text_mask(items, rgb.shape[:2])
    text[95:110, 220:300] = 255
    calls = []
    def reader(path, **kwargs):
        calls.append(kwargs)
        if kwargs.get('recognition_only'):
            with Image.open(path) as image:
                width, height = image.size
            return [{'text': '“点状”', 'box': [0, 0, width, height], 'confidence': .99}], None
        return [{'text': '点状”', 'box': [5, 5, 45, 30], 'confidence': .96}], None
    monkeypatch.setattr('image2editable.layout_first._bright_title_runs',
                        lambda *args: [{'text': '“', 'box': [0, .8, 1, .2]}])
    result, mask, audit = _recover_text_gaps(rgb, items, text, tmp_path, reader)
    assert result[-1]['text'] == '点状”' and 'runs' not in result[-1]
    assert len(calls) == audit['gap_recovery_calls'] == 2
    assert np.all(mask[95:110, 220:300] == 255)


@pytest.mark.parametrize('rows', [2, 3, 5])
def test_quote_retries_share_four_call_budget(tmp_path, monkeypatch, rows):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((rows*180, 600, 3), 255, np.uint8)
    items = []
    for row in range(rows):
        y = row*180+55
        items.append({'text': 'left', 'box': [60, y, 150, 40]})
        cv2.putText(rgb, 'CT', (240, y+30), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    calls = []
    def reader(path, **kwargs):
        calls.append(kwargs)
        if kwargs.get('recognition_only'):
            width, height = Image.open(path).size
            return [{'text': '“点状”', 'box': [0, 0, width, height], 'confidence': .99}], None
        return [{'text': '点状”', 'box': [5, 5, 45, 30], 'confidence': .96}], None
    monkeypatch.setattr('image2editable.layout_first._bright_title_runs', lambda *args: None)
    result, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert len(calls) == audit['gap_recovery_calls'] == 4
    assert len(result) == len(items)+2
    assert sum(bool(call.get('recognition_only')) for call in calls) == 2


@pytest.mark.parametrize('confidence', [.99, .85])
def test_text_gap_scans_short_wrapped_tail_without_lowering_confidence(tmp_path, confidence):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((240, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (60, 145), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': 'previous line', 'box': [60, 40, 420, 40]}]
    calls = []
    def reader(path, **kwargs):
        pixels = np.asarray(Image.open(path).convert('RGB'))
        calls.append(path)
        box = list(cv2.boundingRect((pixels.min(axis=2) < 140).astype(np.uint8)))
        return [{'text': 'CT', 'box': box, 'confidence': confidence}], None
    result, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert len(calls) == 1
    assert audit['gap_texts_recovered'] == int(confidence >= .9)
    if confidence >= .9:
        assert result[-1]['box'][0] >= 60 and result[-1]['box'][1] > 100
    else:
        assert result == items
    blank = np.full_like(rgb, 255)
    _, _, audit = _recover_text_gaps(blank, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert audit['gap_recovery_calls'] == 0


@pytest.mark.parametrize('case', ['covered', 'clipped', 'rule'])
def test_wrapped_tail_scan_rejects_covered_clipped_and_rule_ink(tmp_path, case):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((240, 600, 3), 255, np.uint8)
    items = [{'text': 'previous line', 'box': [60, 40, 420, 40]}]
    if case == 'rule':
        cv2.line(rgb, (50, 130), (170, 130), (30, 30, 30), 3)
    else:
        cv2.putText(rgb, 'CT', (60, 100 if case == 'clipped' else 145),
                    cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    if case == 'covered':
        items.append({'text': 'CT', 'box': [55, 115, 55, 35]})
    calls = []
    def reader(path, **kwargs):
        calls.append(path)
        return [], None
    result, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert result == items and not calls and audit['gap_recovery_calls'] == 0


@pytest.mark.parametrize('distant_neighbor', [False, True])
def test_text_gap_recovers_line_ending_and_skips_blank_tail(tmp_path, distant_neighbor):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': 'left', 'box': [60, 55, 150, 40]}]
    if distant_neighbor:
        items.append({'text': 'far', 'box': [450, 55, 100, 40]})
    calls = []
    def reader(path, **kwargs):
        calls.append(path)
        with Image.open(path) as image:
            pixels = np.asarray(image.convert('RGB'))
        bounds = cv2.boundingRect((pixels.min(axis=2) < 150).astype(np.uint8))
        return [{'text': 'CT', 'box': list(bounds), 'confidence': .99}], None
    result, _, audit = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert len(calls) == 1 and audit['gap_texts_recovered'] == 1
    assert 210 < result[-1]['box'][0] < 300
    blank = np.full_like(rgb, 255)
    unchanged, _, audit = _recover_text_gaps(blank, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert unchanged == items and audit['gap_recovery_calls'] == 0


def test_text_gap_crop_excludes_long_horizontal_rule(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask
    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, 'CT', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    cv2.line(rgb, (300, 75), (380, 75), (30, 30, 30), 3)
    items = [{'text': 'left', 'box': [60, 55, 150, 40]}]
    crops = []
    def reader(path, **kwargs):
        crops.append(np.asarray(Image.open(path)))
        return [], None
    _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert len(crops) == 1
    assert crops[0].shape[1] < 90


def test_source_title_recovery_maps_crop_and_bounds_calls(tmp_path, monkeypatch):
    from image2editable.layout_first import _recover_source_titles

    monkeypatch.setattr('scripts.art_text.estimate_art_text_runs', lambda *args, **kwargs: [{'text': 'TITLE'}])

    rgb = np.full((600, 1000, 3), 230, np.uint8)
    cv2.putText(rgb, 'TITLE', (480, 250), cv2.FONT_HERSHEY_SIMPLEX, 2, (20, 20, 20), 4)
    items = [{'text': 'subtitle', 'box': [480, 90, 300, 30]},
             {'text': 'button', 'box': [480, 330, 350, 40]}]
    calls = []
    def reader(path, **kwargs):
        calls.append(kwargs)
        return [{'text': 'TITLE', 'box': [30, 50, 300, 80], 'confidence': .99}], None
    recovered, audit = _recover_source_titles(rgb, items, tmp_path, reader)
    assert len(calls) == 1 and calls[0]['recognition_only']
    assert recovered[-1]['text'] == 'TITLE'
    assert recovered[-1]['box'][0] > 400 and recovered[-1]['box'][1] > 120
    assert audit['titles_recovered'] == 1
    unchanged, audit = _recover_source_titles(rgb, items, tmp_path, lambda *a, **k: ([{'text': 'guess', 'box': [0, 0, 300, 80], 'confidence': .6}], None))
    assert unchanged == items and audit['titles_recovered'] == 0
    monkeypatch.setattr('scripts.art_text.estimate_art_text_runs', lambda *args, **kwargs: None)
    unchanged, audit = _recover_source_titles(rgb, items, tmp_path, reader)
    assert unchanged == items
    assert audit['titles_recovered'] == 0
    assert audit['title_candidates_preserved'][0]['text'] == 'TITLE'


def test_source_title_recovery_accepts_bright_title_on_dark_backdrop(tmp_path, monkeypatch):
    from image2editable.layout_first import _recover_source_titles

    monkeypatch.setattr('scripts.art_text.estimate_art_text_runs', lambda *args, **kwargs: [{'text': 'TITLE'}])

    rgb = np.full((600, 1000, 3), (20, 60, 160), np.uint8)
    cv2.putText(rgb, 'TITLE', (480, 250), cv2.FONT_HERSHEY_SIMPLEX, 2, (255, 255, 255), 8)
    items = [{'text': 'subtitle', 'box': [480, 90, 300, 30]},
             {'text': 'button', 'box': [480, 330, 350, 40]}]
    calls = []

    def reader(path, **kwargs):
        calls.append(kwargs)
        return [{'text': 'TITLE', 'box': [30, 50, 300, 80], 'confidence': .99}], None

    recovered, audit = _recover_source_titles(rgb, items, tmp_path, reader)
    assert len(calls) == 1
    assert audit['titles_recovered'] == 1


def test_bright_title_keeps_letter_spacing_and_rejects_bad_fit(monkeypatch):
    from image2editable.layout_first import _bright_title_runs
    pixels = np.full((160, 400, 3), 80, np.uint8)
    cv2.putText(pixels, 'A', (25, 120), cv2.FONT_HERSHEY_SIMPLEX, 3, (255, 255, 255), 5)
    cv2.putText(pixels, 'B', (240, 120), cv2.FONT_HERSHEY_SIMPLEX, 3, (255, 255, 255), 5)
    item = {'text': 'AB', 'box': [0, 0, 400, 160], 'font': 'Arial'}
    monkeypatch.setattr('scripts.font_match.match_glyph', lambda *a: {'font': 'Arial', 'bold': True, 'rotation': 0, 'font_size': 100, 'fit_iou': .96})
    runs = _bright_title_runs(pixels, item, 960)
    assert [run['text'] for run in runs] == ['A', 'B']
    assert runs[1]['box'][0] > .5
    assert all(run['color'] == '#ffffff' for run in runs)
    monkeypatch.setattr('scripts.font_match.match_glyph', lambda *a: {'fit_iou': .6})
    assert _bright_title_runs(pixels, item, 960) is None


def test_positioned_title_mask_preserves_space_between_letters():
    from image2editable.layout_first import _text_mask
    item = {'box': [20, 20, 300, 120], 'runs': [
        {'box': [.1, .1, .2, .6]}, {'box': [.7, .1, .2, .6]}]}
    mask = _text_mask([item], (180, 360))
    assert mask[60, 60] == 255
    assert mask[60, 170] == 0
    assert mask[25, 25] == 0


def test_layout_group_text_keeps_fixed_frame(tmp_path):
    from scripts.ppt_assemble import assemble_pptx_multi
    from pptx.enum.text import MSO_AUTO_SIZE

    output = tmp_path / 'group.pptx'
    assemble_pptx_multi([{
        'img_width': 640, 'img_height': 360, 'background_rgb': [255, 255, 255],
        'components': [], 'visual_elements': [],
        'text_items': [{'text': 'A long editable sentence', 'box': [100, 100, 400, 30],
                        'font': 'Arial', 'font_size': 16, 'layout_object_id': 'text_0001'}],
        'object_groups': {'card_group_0001': ['text_0001']},
    }], output, add_reference=False)
    text = Presentation(output).slides[0].shapes[0].shapes[0]
    assert text.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert text.text_frame.word_wrap is False


def _source(tmp_path):
    from image2editable.shape_analysis import _rounded_rectangle

    pixels = np.full((360, 640, 3), 255, np.uint8)
    body = _rounded_rectangle(120, 220, 16)
    pixels[100:220, 60:280][body] = (20, 100, 220)
    pixels[100:220, 360:580][body] = (200, 60, 50)
    cv2.rectangle(pixels, (110, 140), (230, 165), (255, 255, 255), -1)
    cv2.line(pixels, (390, 180), (420, 150), (15, 15, 15), 7)
    cv2.line(pixels, (420, 150), (450, 180), (15, 15, 15), 7)
    path = tmp_path / 'source.png'
    Image.fromarray(pixels).save(path)
    text = [{'box': [110, 140, 121, 26], 'text': 'CARD', 'font_size': 16,
             'font': 'Arial', 'color': '#ffffff', 'bold': True}]
    return path, text


def test_cli_exposes_layout_first_without_changing_legacy_defaults():
    args = cli.build_parser().parse_args(
        ['convert', 'page.png', '--layout-first', '--card-mode', 'png'])
    assert args.layout_first
    assert args.card_mode == 'png'
    assert not cli.build_parser().parse_args(['convert', 'page.png']).layout_first


def test_native_cards_are_independent_and_not_baked_in_background(tmp_path):
    from image2editable.layout_first import prepare_layout

    path, text = _source(tmp_path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=text)
    assert slide['background_rgb'] == [255, 255, 255]
    cards = [e for e in slide['visual_elements'] if e['route'] == 'native_shape']
    assert len(cards) == 2
    assert all(e['shape']['shape_type'] == 'rounded_rectangle' for e in cards)
    assert len(slide['text_items']) == 1
    pictures = [e for e in slide['visual_elements'] if e['route'] == 'raster_component']
    assert len(pictures) == 1
    assert pictures[0]['component']['w'] < 100
    assert audit['native_shapes'] == 2
    assert audit['residual_objects'] == 1
    assert len(slide['object_groups']) == 2
    assert any('object_0001' in ids for ids in slide['object_groups'].values())


def test_png_cards_keep_complete_outline_and_remove_native_text_pixels(tmp_path):
    from image2editable.layout_first import prepare_layout

    path, text = _source(tmp_path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=text, card_mode='png')
    assert audit['native_shapes'] == 0
    assert audit['card_images'] == 2
    cards = [e for e in slide['visual_elements'] if e.get('role') == 'card']
    assert len(cards) == 2
    with Image.open(cards[0]['component']['path']) as image:
        rgba = np.asarray(image)
    assert rgba.shape[:2] == (120, 220)
    assert rgba[60, 100, 3] == 255
    assert np.max(np.abs(rgba[45, 100, :3].astype(int) - [20, 100, 220])) <= 2
    assert rgba[0, 0, 3] == 0
    assert slide['background_rgb'] == [255, 255, 255]


def test_text_cache_is_bound_to_source_not_arbitrary_previous_page(tmp_path):
    from image2editable.layout_first import load_text_cache

    path, text = _source(tmp_path)
    cache = tmp_path / 'text.json'
    cache.write_text(json.dumps({'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                 'text_items': text}), encoding='utf-8')
    assert load_text_cache(path, cache) == text
    cache.write_text(json.dumps({'source_sha256': '0' * 64, 'text_items': text}), encoding='utf-8')
    with pytest.raises(ValueError, match='source'):
        load_text_cache(path, cache)


def test_convert_is_real_pptx_with_native_shapes_text_and_icon(tmp_path):
    from image2editable.layout_first import convert_layout

    path, text = _source(tmp_path)
    cache = tmp_path / 'text.json'
    cache.write_text(json.dumps({'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                 'text_items': text}), encoding='utf-8')
    output = tmp_path / 'editable.pptx'
    report = convert_layout([path], output_path=output, run_dir=tmp_path / 'run',
                            text_cache=cache, slide_size='16:9', render=False)
    prs = Presentation(output)
    assert len(prs.slides) == 1
    shapes = prs.slides[0].shapes
    assert sum(s.shape_type == 6 for s in shapes) == 2
    children = [child for s in shapes if s.shape_type == 6 for child in s.shapes]
    assert sum(s.shape_type == 1 for s in children) == 2
    assert sum(s.shape_type == 13 for s in children) == 1
    assert sum(s.has_text_frame and bool(s.text) for s in children) == 1
    assert report['pages'][0]['native_shapes'] == 2
    assert all(s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style') is None
               for s in shapes if s.shape_type == 1)
    assert report['pages'][0]['render_status'] == 'not_attempted'
    with pytest.raises(FileExistsError):
        convert_layout([path], output_path=output, run_dir=tmp_path / 'other', render=False)


def test_texture_is_preserved_as_one_object_not_claimed_as_native_shape(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    rng = np.random.default_rng(2)
    pixels[40:160, 50:270] = rng.integers(20, 200, (120, 220, 3), np.uint8)
    path = tmp_path / 'photo.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert audit['native_shapes'] == 0
    assert audit['residual_objects'] == 1
    assert len(slide['visual_elements']) == 1


def test_outline_frame_is_reconstructed_instead_of_kept_in_background(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    cv2.rectangle(pixels, (20, 30), (300, 180), (218, 228, 244), 2)
    path = tmp_path / 'frame.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert audit['native_shapes'] == 1
    shape = slide['visual_elements'][0]['shape']
    assert shape['fill_rgb'] is None
    assert shape['stroke_rgb'] is not None
    assert slide['background_rgb'] == [255, 255, 255]


def test_dashed_separator_inside_outline_remains_a_native_line(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    cv2.rectangle(pixels, (20, 30), (300, 180), (218, 228, 244), 2)
    for x in range(45, 280, 10):
        cv2.line(pixels, (x, 100), (x+5, 100), (208, 220, 240), 2)
    path = tmp_path / 'lines.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    lines = [e for e in slide['visual_elements'] if e.get('role') == 'line']
    assert len(lines) == 1
    assert lines[0]['shape']['dash'] == 'dash'


def test_pale_card_is_filled_native_body_not_outline(tmp_path):
    from image2editable.layout_first import prepare_layout
    from image2editable.shape_analysis import _rounded_rectangle

    pixels = np.full((200, 320, 3), 255, np.uint8)
    body = _rounded_rectangle(100, 220, 12)
    pixels[50:150, 50:270][body] = (232, 239, 252)
    path = tmp_path / 'pale.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert audit['native_shapes'] == 1
    shape = slide['visual_elements'][0]['shape']
    assert shape['fill_rgb'] == [232, 239, 252]
    assert shape['stroke_rgb'] is None


def test_pale_frame_stroke_uses_ink_not_white_antialias_pixels(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    cv2.rectangle(pixels, (20, 30), (300, 180), (218, 228, 244), 2, cv2.LINE_AA)
    path = tmp_path / 'frame-aa.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    shape = slide['visual_elements'][0]['shape']
    assert max(abs(a-b) for a, b in zip(shape['stroke_rgb'], [218, 228, 244])) <= 5


def test_uncached_convert_creates_ocr_worker_root_and_records_empty_warning(tmp_path, monkeypatch):
    from image2editable.layout_first import convert_layout
    from scripts import text_detect

    path, _ = _source(tmp_path)
    def detect(source, **kwargs):
        assert kwargs['worker_root'].is_dir()
        return [], np.zeros((360, 640), np.uint8)
    monkeypatch.setattr(text_detect, 'detect_text', detect)
    report = convert_layout([path], output_path=tmp_path / 'page.pptx', run_dir=tmp_path / 'run', render=False)
    assert any('OCR' in warning for warning in report['pages'][0]['warnings'])


def test_header_and_body_become_one_complete_png_card(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((240, 400, 3), 255, np.uint8)
    pixels[50:90, 100:300] = [20, 100, 220]
    pixels[90:190, 100:300] = [232, 239, 252]
    path = tmp_path / 'compound.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='png')
    cards = [e for e in slide['visual_elements'] if e.get('role') == 'card']
    assert len(cards) == 1
    assert cards[0]['component']['w'] == 200
    assert cards[0]['component']['h'] == 140
    assert audit['card_images'] == 1


def test_png_outline_uses_full_resolution_not_upsampled_detection_mask(tmp_path):
    from image2editable.layout_first import prepare_layout
    from image2editable.shape_analysis import _rounded_rectangle

    pixels = np.full((1080, 1920, 3), 255, np.uint8)
    mask = _rounded_rectangle(157, 321, 23)
    pixels[213:370, 137:458][mask] = (20, 100, 220)
    path = tmp_path / 'large.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='png')
    card = next(e for e in slide['visual_elements'] if e['role'] == 'card')
    with Image.open(card['component']['path']) as image:
        alpha = np.asarray(image)[:, :, 3]
    assert alpha.shape == mask.shape
    assert np.array_equal(alpha != 0, mask)


def test_low_saturation_card_survives_text_box_touching_edge(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    pixels[30:180, 20:300] = [236, 241, 249]
    text = [{'box': [20, 65, 280, 30], 'text': 'FULL WIDTH', 'font_size': 14, 'font': 'Arial'}]
    cv2.rectangle(pixels, (80, 72), (210, 88), (20, 20, 20), -1)
    path = tmp_path / 'pale-text.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=text)
    assert audit['cards'] == 1
    assert slide['visual_elements'][0]['shape']['fill_rgb'] == [236, 241, 249]


def test_raster_arrow_text_cleanup_keeps_opaque_background_under_native_text(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    points = np.array([[20, 40], [230, 40], [290, 100], [230, 160], [20, 160], [80, 100]], np.int32)
    cv2.fillPoly(pixels, [points], [20, 100, 220])
    cv2.rectangle(pixels, (120, 80), (200, 120), (255, 255, 255), -1)
    text = [{'box': [120, 80, 81, 41], 'text': 'ARROW', 'font_size': 14, 'font': 'Arial'}]
    path = tmp_path / 'arrow.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=text)
    picture = next(e for e in slide['visual_elements'] if e['route'] == 'raster_component')
    c = picture['component']
    with Image.open(c['path']) as image:
        pixel = np.asarray(image)[100-c['y'], 160-c['x']]
    assert pixel[3] == 255
    assert np.max(np.abs(pixel[:3].astype(int)-[20, 100, 220])) <= 4


def test_cli_routes_layout_conversion_to_new_backend(tmp_path, monkeypatch, capsys):
    from image2editable import layout_first

    calls = []
    def convert(sources, **kwargs):
        calls.append((sources, kwargs))
        return {'status': 'assembled'}
    monkeypatch.setattr(layout_first, 'convert_layout', convert)
    assert cli.main(['convert', 'source.png', '--layout-first', '--slide-size', 'original',
                     '--run-dir', str(tmp_path / 'run'), '-o', str(tmp_path / 'out.pptx')]) == 0
    assert calls[0][1]['slide_size'] == 'original'
    assert calls[0][1]['card_mode'] == 'native'
    assert json.loads(capsys.readouterr().out)['status'] == 'assembled'


def test_layout_conversion_does_not_silently_ignore_both_slide_sizes(tmp_path):
    from image2editable.layout_first import convert_layout

    path, _ = _source(tmp_path)
    with pytest.raises(ValueError, match='slide-size'):
        convert_layout([path], output_path=tmp_path / 'out.pptx', run_dir=tmp_path / 'run', slide_size='both')
    assert not (tmp_path / 'run').exists()


def test_colored_outline_and_pale_fill_do_not_create_two_overlapping_cards(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    cv2.rectangle(pixels, (20, 40), (300, 160), [255, 240, 240], -1)
    cv2.rectangle(pixels, (20, 40), (300, 160), [245, 140, 140], 3)
    path = tmp_path / 'border-fill.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert audit['cards'] == 1
    assert len(slide['object_groups']) == 1


def test_photo_keeps_white_boundary_pixels_in_one_complete_object(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((300, 400, 3), 255, np.uint8)
    pixels[40:260, 50:350] = np.random.default_rng(3).integers(20, 200, (220, 300, 3), np.uint8)
    cv2.fillPoly(pixels, [np.array([[50, 100], [90, 140], [50, 180]])], [255, 255, 255])
    path = tmp_path / 'white-coat.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    picture = next(e for e in slide['visual_elements'] if e['route'] == 'raster_component')
    c = picture['component']
    with Image.open(c['path']) as image:
        pixel = np.asarray(image)[140-c['y'], 55-c['x']]
    assert pixel.tolist() == [255, 255, 255, 255]


def test_render_failure_is_recorded_without_losing_assembled_report(tmp_path, monkeypatch):
    from image2editable.layout_first import convert_layout
    from image2editable.powerpoint_renderer import PowerPointRenderer

    path, text = _source(tmp_path)
    cache = tmp_path / 'text.json'
    cache.write_text(json.dumps({'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                 'text_items': text}), encoding='utf-8')
    class BrokenRenderer:
        def render_page(self, *args, **kwargs):
            raise AttributeError('COM disconnected')
    monkeypatch.setattr(PowerPointRenderer, 'discover', lambda: BrokenRenderer())
    report = convert_layout([path], output_path=tmp_path / 'page.pptx', run_dir=tmp_path / 'run', text_cache=cache)
    assert report['pages'][0]['render_status'] == 'failed'
    assert (tmp_path / 'run' / 'layout-report.json').is_file()
    assert (tmp_path / 'page.pptx').is_file()


def test_wide_text_mask_does_not_split_arrow_into_top_and_bottom_fragments(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((200, 320, 3), 255, np.uint8)
    points = np.array([[20, 40], [230, 40], [290, 100], [230, 160], [20, 160], [80, 100]], np.int32)
    cv2.fillPoly(pixels, [points], [20, 100, 220])
    cv2.rectangle(pixels, (120, 80), (200, 120), (255, 255, 255), -1)
    text = [{'box': [20, 75, 270, 50], 'text': 'ARROW', 'font_size': 14, 'font': 'Arial'}]
    path = tmp_path / 'wide-arrow.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=text)
    pictures = [e for e in slide['visual_elements'] if e['route'] == 'raster_component']
    assert len(pictures) == 1
    c = pictures[0]['component']
    with Image.open(c['path']) as image:
        pixel = np.asarray(image)[100-c['y'], 160-c['x']]
    assert pixel[3] == 255


def test_batch_reuses_one_bound_text_cache_per_source(tmp_path):
    from image2editable.layout_first import convert_layout

    path, text = _source(tmp_path)
    cache = tmp_path / 'text.json'
    cache.write_text(json.dumps({'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                 'text_items': text}), encoding='utf-8')
    report = convert_layout([path, path], output_path=tmp_path / 'two.pptx', run_dir=tmp_path / 'run',
                            text_cache=[cache, cache], render=False)
    assert len(Presentation(tmp_path / 'two.pptx').slides) == 2
    assert all(p['ocr_reused'] and p['native_shapes'] == 2 for p in report['pages'])


def test_irregular_raster_header_joins_pale_body_as_one_complete_png(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((260, 360, 3), 255, np.uint8)
    pixels[100:230, 60:300] = [232, 239, 252]
    points = np.array([[60, 50], [270, 50], [300, 75], [270, 100], [60, 100], [90, 75]])
    cv2.fillPoly(pixels, [points], [20, 100, 220])
    path = tmp_path / 'raster-head.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='png')
    cards = [e for e in slide['visual_elements'] if e['role'] == 'card']
    assert len(cards) == 1
    assert cards[0]['component']['h'] >= 180
    assert audit['residual_objects'] == 0


def test_shared_suffix_prefix_glyph_is_not_printed_twice():
    from image2editable.layout_first import refine_layout_text

    pixels = np.full((150, 320, 3), 255, np.uint8)
    first = {'box': [10, 40, 110, 50], 'text': 'Hi X', 'color': '#0066cc', 'confidence': .95,
             'words': [{'text': 'Hi', 'box': [0, 0, .6, 1]}, {'text': 'X', 'box': [.84, 0, .16, 1]}]}
    second = {'box': [100, 40, 100, 50], 'text': 'X Y', 'color': '#222222', 'confidence': .99,
              'words': [{'text': 'X', 'box': [0, 0, .4, 1]}, {'text': 'Y', 'box': [.5, 0, .5, 1]}]}
    items, audit = refine_layout_text(pixels, [first, second])
    assert [i['text'] for i in items] == ['Hi', 'X Y']
    assert audit['duplicate_spans_removed'] == 1
    assert first['text'] == 'Hi X'


def test_repeated_word_at_different_positions_is_preserved():
    from image2editable.layout_first import refine_layout_text

    pixels = np.full((150, 320, 3), 255, np.uint8)
    words = [{'text': 'X', 'box': [0, 0, 1, 1]}]
    first = {'box': [10, 40, 40, 50], 'text': 'X', 'words': words}
    second = {'box': [200, 40, 40, 50], 'text': 'X', 'words': words}
    items, audit = refine_layout_text(pixels, [first, second])
    assert len(items) == 2
    assert audit['duplicate_spans_removed'] == 0


def test_mixed_text_colors_are_preserved_without_changing_words():
    from image2editable.layout_first import refine_layout_text

    pixels = np.full((100, 200, 3), 255, np.uint8)
    pixels[32:68, 30:55] = [0, 102, 204]
    pixels[32:68, 100:125] = [32, 32, 32]
    item = {'box': [20, 25, 120, 50], 'text': 'AB', 'color': '#222222', 'font_size': 18,
            'words': [{'text': 'A', 'box': [0, 0, .4, 1]}, {'text': 'B', 'box': [.6, 0, .4, 1]}]}
    items, audit = refine_layout_text(pixels, [item])
    assert ''.join(i['text'] for i in items) == 'AB'
    assert [i['color'] for i in items] == ['#0066cc', '#202020']
    assert audit['mixed_color_lines'] == 1


def test_uncovered_label_recovery_is_bounded_and_does_not_scan_plain_cards(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((360, 640, 3), 255, np.uint8)
    pixels[100:220, 60:280] = [20, 100, 220]
    for x in (110, 140, 170):
        pixels[145:170, x:x+12] = 255
    pixels[100:220, 360:580] = [200, 60, 50]
    source = tmp_path / 'labels.png'
    Image.fromarray(pixels).save(source)
    calls = []
    def read_label(path, **kwargs):
        calls.append((path, kwargs))
        return [{'box': [50, 45, 72, 25], 'text': 'TAG', 'font_size': 14, 'confidence': .99}], None
    slide, audit = prepare_layout(source, tmp_path / 'assets', text_items=[], label_reader=read_label)
    assert len(calls) == 1
    assert len(slide['text_items']) == 1
    assert audit['text_audit']['label_recovery_calls'] == 1
    assert audit['text_audit']['labels_recovered'] == 1


def test_rotation_cleanup_mask_covers_rotated_label_corners():
    from image2editable.layout_first import _text_mask

    item = {'box': [40, 40, 80, 20], 'text': 'LABEL', 'rotation_degrees': 20}
    mask = _text_mask([item], (120, 160))
    assert mask[28:36, 40:65].any()
    assert mask[66:76, 100:130].any()


def test_batch_ocr_reuses_one_task_worker_and_closes_it(tmp_path, monkeypatch):
    from image2editable.layout_first import convert_layout
    from scripts import text_detect

    source, text = _source(tmp_path)
    pools = []
    def detect(path, **kwargs):
        pools.append(kwargs.get('worker_pool'))
        return text, np.zeros((360, 640), np.uint8)
    monkeypatch.setattr(text_detect, 'detect_text', detect)
    convert_layout([source, source], output_path=tmp_path / 'batch.pptx', run_dir=tmp_path / 'run', render=False)
    assert len(pools) == 2
    assert pools[0] is not None and pools[0] is pools[1]
    assert pools[0]._closed


def test_truncated_duplicate_prefix_is_removed_from_weaker_detection():
    from image2editable.layout_first import refine_layout_text

    pixels = np.full((150, 320, 3), 255, np.uint8)
    first = {'box': [10, 40, 130, 50], 'text': 'Hi X', 'confidence': .999,
             'words': [{'text': 'Hi', 'box': [0, 0, .6, 1]}, {'text': 'X', 'box': [.7, 0, .3, 1]}]}
    second = {'box': [115, 40, 100, 50], 'text': 'X Y', 'confidence': .8,
              'words': [{'text': 'X', 'box': [0, 0, .15, 1]}, {'text': 'Y', 'box': [.4, 0, .6, 1]}]}
    items, audit = refine_layout_text(pixels, [first, second])
    assert [i['text'] for i in items] == ['Hi X', 'Y']
    assert audit['duplicate_spans_removed'] == 1


def test_narrow_punctuation_inherits_adjacent_text_color():
    from image2editable.layout_first import refine_layout_text

    pixels = np.full((100, 200, 3), 255, np.uint8)
    pixels[32:68, 30:55] = [0, 102, 204]
    pixels[32:68, 110:135] = [32, 32, 32]
    item = {'box': [20, 25, 140, 50], 'text': 'A:B', 'color': '#202020', 'font_size': 18,
            'words': [{'text': 'A', 'box': [0, 0, .35, 1]}, {'text': ':', 'box': [.36, 0, .01, 1]},
                      {'text': 'B', 'box': [.6, 0, .35, 1]}]}
    items, _ = refine_layout_text(pixels, [item])
    assert items[0]['text'] == 'A:'
    assert items[0]['color'] == '#0066cc'
    assert ''.join(i['text'] for i in items) == 'A:B'


def test_source_backdrop_cleanup_keeps_dark_surface_under_bright_title():
    from image2editable.layout_first import _clean_source_backdrop

    rgb = np.full((400, 800, 3), (30, 90, 200), np.uint8)
    rgb[160:172, 110:690] = (240, 245, 250)
    affected = np.zeros((400, 800), np.uint8)
    affected[172:322, 120:680] = 255
    rgb[172:322, 120:680] = (245, 245, 245)
    cleaned = _clean_source_backdrop(rgb, affected)
    center = cleaned[247, 400]
    assert center.min() < 220 and abs(int(center[2])-200) < 40


def test_source_background_cleanup_preserves_smooth_gradient_without_panel():
    from image2editable.layout_first import _clean_source_backdrop

    yy, xx = np.indices((160, 320))
    curve = 32*np.exp(-((xx-165)**2+(yy-75)**2)/1600)
    original = np.stack((180+xx*.1+yy*.15+curve, 200+xx*.08+curve, 220+yy*.1+curve*.4), axis=2).clip(0, 255).astype(np.uint8)
    rgb = original.copy()
    affected = np.zeros((160, 320), np.uint8)
    affected[60:95, 120:220] = 255
    cv2.putText(rgb, 'TEXT', (122, 84), cv2.FONT_HERSHEY_SIMPLEX, .7, (0, 0, 0), 2)
    cleaned = _clean_source_backdrop(rgb, affected)
    assert np.array_equal(cleaned[affected == 0], rgb[affected == 0])
    assert np.percentile(np.abs(cleaned[affected != 0].astype(int)-original[affected != 0]), 95) <= 6


def test_saturated_glyph_with_enclosed_counter_is_not_a_card(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((400, 800, 3), 255, np.uint8)
    cv2.ellipse(pixels, (200, 200), (47, 51), 0, 0, 360, (30, 90, 220), -1)
    cv2.ellipse(pixels, (200, 200), (20, 24), 0, 0, 360, (255, 255, 255), -1)
    cv2.rectangle(pixels, (450, 120), (700, 280), (30, 90, 220), -1)
    path = tmp_path / 'glyph.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='native')
    cards = [o for o in audit['objects'] if o['role'] == 'card']
    assert len(cards) == 1 and cards[0]['bbox'][0] > 400


def test_in_card_icon_halo_extracts_with_soft_alpha(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((500, 900, 3), 255, np.uint8)
    cv2.rectangle(pixels, (200, 100), (700, 400), (30, 110, 230), -1)
    glow = np.zeros((500, 900), np.float32)
    cv2.circle(glow, (450, 250), 110, 1.0, -1)
    glow = cv2.GaussianBlur(glow, (0, 0), 40)
    halo = np.stack([np.full_like(glow, 200), np.full_like(glow, 220), np.full_like(glow, 255)], axis=2)
    pixels = np.round(pixels.astype(np.float32) * (1 - glow[..., None] * .35)
                      + halo * glow[..., None] * .35).astype(np.uint8)
    cv2.circle(pixels, (450, 250), 45, (250, 180, 40), -1)
    path = tmp_path / 'halo.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='native')
    visuals = [e for e in slide['visual_elements'] if e.get('role') == 'visual']
    assert visuals
    alpha = np.asarray(Image.open(visuals[0]['component']['path']))[..., 3]
    assert np.count_nonzero((alpha > 32) & (alpha < 255)) > 0


def test_readable_glyph_in_unreadable_art_unit_is_suppressed(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1000, 3), 255, np.uint8)
    cv2.rectangle(pixels, (150, 100), (850, 500), (252, 248, 240), -1)
    cv2.rectangle(pixels, (300, 180), (420, 380), (143, 191, 249), -1)
    cv2.rectangle(pixels, (480, 180), (560, 380), (143, 191, 249), -1)
    cv2.circle(pixels, (470, 450), 35, (58, 155, 250), -1)
    items = [{'text': '0', 'box': [298, 178, 124, 204], 'color': '#8fbff9', 'font_size': 80, 'confidence': .95},
             {'text': '标题', 'box': [200, 115, 300, 50], 'color': '#0049dc', 'font_size': 40, 'confidence': .95}]
    path = tmp_path / 'unit.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items, card_mode='native')
    assert '0' not in [i['text'] for i in slide['text_items']]
    assert '标题' in [i['text'] for i in slide['text_items']]


def test_glow_bridged_owned_components_split_into_separate_objects(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1000, 3), 255, np.uint8)
    cv2.rectangle(pixels, (150, 60), (850, 560), (252, 248, 240), -1)
    cv2.fillConvexPoly(pixels, np.array([[(350, 250), (480, 250), (415, 120)]]), (250, 155, 58))
    cv2.fillConvexPoly(pixels, np.array([[(350, 380), (480, 380), (415, 510)]]), (250, 155, 58))
    cv2.rectangle(pixels, (390, 250), (440, 380), (238, 224, 208), -1)
    path = tmp_path / 'bridge.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='native')
    covering = [e for e in slide['visual_elements'] if e.get('role') == 'visual'
                and e['component']['x'] < 500 and e['component']['y'] < 300
                and e['component']['x']+e['component']['w'] > 350]
    assert covering
    assert not any(e['component']['y'] < 300 and e['component']['y']+e['component']['h'] > 400
                   for e in covering)


def test_flatbg_illustration_does_not_lose_card_cut_columns(tmp_path):
    from image2editable.layout_first import prepare_layout

    source, items = _bany06_fixture('flatbg', 2, 'b')
    rgb = np.asarray(Image.open(source).convert('RGB')).copy()
    path = tmp_path/'real-crop.png'
    Image.fromarray(rgb).save(path)
    slide, _ = prepare_layout(path, tmp_path/'assets', text_items=items, background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    for x, y in ((775, 979), (775, 1677)):
        covered = False
        for element in slide['visual_elements']:
            c = element.get('component')
            if c and c['x'] <= x < c['x']+c['w'] and c['y'] <= y < c['y']+c['h']:
                rgba = np.asarray(Image.open(c['path']).convert('RGBA'))
                covered |= rgba[y-c['y'], x-c['x'], 3] > 0
        assert covered, (x, y)
    assert not any(e.get('component') and e['component']['w']*e['component']['h'] < rgb.shape[0]*rgb.shape[1]*.0001
                   for e in slide['visual_elements'])


def test_illustration_beside_card_is_not_cut_at_card_left(tmp_path):
    from image2editable.layout_first import prepare_layout

    rgb = np.full((600, 1000, 3), 255, np.uint8)
    cv2.ellipse(rgb, (350, 400), (300, 190), 0, 0, 360, (180, 80, 70), -1)
    cv2.ellipse(rgb, (350, 400), (240, 130), 0, 0, 360, (70, 90, 180), -1)
    cv2.rectangle(rgb, (350, 35), (950, 180), (220, 235, 250), -1)
    path = tmp_path/'illustration.png'
    Image.fromarray(rgb).save(path)
    slide, _ = prepare_layout(path, tmp_path/'assets', text_items=[], label_reader=lambda *a, **k: ([], None))
    visuals = [e['component'] for e in slide['visual_elements'] if e.get('component')]
    assert any(c['x'] < 348 and c['x']+c['w'] > 353 and c['y'] <= 400 < c['y']+c['h'] for c in visuals)


def test_micro_free_residual_does_not_emit_independent_object(tmp_path):
    from image2editable.layout_first import _emit_residual

    rgba = np.full((22, 15, 4), 255, np.uint8)
    elements = []
    count = _emit_residual([{'rgba': rgba, 'x': 50, 'y': 50, 'w': 15, 'h': 22, 'owned': False}],
                           elements, tmp_path, np.full((2160, 3840), 100, np.uint8),
                           (2160, 3840), np.zeros((2160, 3840), np.uint8), [])
    assert count == 0


def test_label_size_uses_measured_glyph_height(monkeypatch):
    from image2editable.layout_first import refine_layout_text
    from scripts.font_match import resolve_font
    from PIL import ImageDraw

    image = Image.new('RGB', (1200, 220), 'white')
    font = resolve_font('Microsoft YaHei', True, size=80)
    ImageDraw.Draw(image).text((30, 20), '课堂提问：', font=font, fill=(0, 55, 182))
    ImageDraw.Draw(image).text((500, 20), '正文', font=font, fill=(50, 50, 50))
    monkeypatch.setattr('scripts.font_match.match_text_face', lambda *a, **k:
                        {'font': 'Microsoft YaHei', 'bold': True, 'font_size_px': 30})
    item = {'text': '课堂提问：正文', 'box': [20, 10, 700, 140], 'font_size': 14,
            'words': [{'text': '课堂提问：', 'box': [0, 0, .65, 1]},
                      {'text': '正文', 'box': [.68, 0, .32, 1]}]}
    result, _ = refine_layout_text(np.asarray(image), [item])
    assert abs(result[0]['font_size'] - 80*960/1200) < 2
    assert result[0]['box_kind'] == 'ink'


def test_round_corner_fit_uses_whole_profile_not_first_row():
    from image2editable.layout_first import _corner_radius, _rounded_rectangle

    mask = _rounded_rectangle(600, 1500, 85).astype(np.uint8)
    mask[0, 70:1430] = 1
    assert abs(_corner_radius(mask)-85) <= 1


def test_recovered_fullwidth_punctuation_preserves_cjk_advance(tmp_path):
    from image2editable.layout_first import _recover_text_gaps, _text_mask

    rgb = np.full((160, 600, 3), 255, np.uint8)
    cv2.putText(rgb, '?', (240, 85), cv2.FONT_HERSHEY_SIMPLEX, 1, (30, 30, 30), 2)
    items = [{'text': '癌', 'box': [60, 55, 150, 40],
              'words': [{'text': '癌', 'box': [0, 0, 1, 1]}]}]
    def reader(path, **kwargs):
        return [{'text': '癌?', 'box': [50, 5, 150, 30], 'confidence': .99,
                 'words': [{'text': '癌', 'box': [0, 0, .5, 1]},
                           {'text': '?', 'box': [.5, 0, .5, 1]}]}], None
    result, _, _ = _recover_text_gaps(rgb, items, _text_mask(items, rgb.shape[:2]), tmp_path, reader)
    assert result[0]['text'] == '癌？'


def test_card_edge_sliver_is_not_an_independent_unit(tmp_path):
    from image2editable.layout_first import _dedupe_card_rasters

    rgba=np.full((190,4,4),255,np.uint8);rgba[:,:,:3]=[30,60,140]
    path=tmp_path/'edge.png';Image.fromarray(rgba).save(path)
    elements=[{'object_id':'card','role':'card','route':'native_shape','bbox':[0,0,440,290],
               'shape':{'fill_rgb':[30,60,140]}},
              {'object_id':'unit','role':'visual','route':'raster_component',
               'component':{'path':str(path),'x':437,'y':30,'w':4,'h':190}}]
    result,_=_dedupe_card_rasters(elements)
    assert [e['object_id'] for e in result]==['card']


def test_tail_punctuation_context_merges_without_repeating_glyph(tmp_path):
    from image2editable.layout_first import _recover_text_gaps,_text_mask

    rgb=np.full((160,600,3),255,np.uint8)
    cv2.putText(rgb,'?',(240,85),cv2.FONT_HERSHEY_SIMPLEX,1,(30,30,30),2)
    items=[{'text':'癌','box':[60,55,150,40]}]
    def reader(path,**kwargs):
        return [{'text':'癌?','box':[50,5,60,30],'confidence':.99}],None
    result,_,_=_recover_text_gaps(rgb,items,_text_mask(items,rgb.shape[:2]),tmp_path,reader)
    assert [v['text'] for v in result]==['癌？']


def test_tiny_owned_residual_does_not_emit_independent_object(tmp_path):
    from image2editable.layout_first import _emit_residual

    rgba=np.full((21,7,4),255,np.uint8)
    elements=[]
    count=_emit_residual([{'rgba':rgba,'x':50,'y':50,'w':7,'h':21,'owned':True}],
                         elements,tmp_path,np.full((2160,3840),100,np.uint8),
                         (2160,3840),np.zeros((2160,3840),np.uint8),[])
    assert count==0


def test_split_color_label_refits_its_own_weight(monkeypatch):
    from image2editable.layout_first import refine_layout_text

    rgb=np.full((100,300,3),255,np.uint8)
    cv2.rectangle(rgb,(20,20),(80,70),(20,60,180),-1)
    cv2.rectangle(rgb,(100,20),(220,70),(50,50,50),-1)
    monkeypatch.setattr('scripts.font_match.match_text_face',lambda *a,**k:
                        {'font':'Microsoft YaHei','bold':True,'font_size_px':80})
    item={'text':'标签：正文','box':[10,10,230,70],'bold':False,'font_size':14,
          'words':[{'text':'标签：','box':[0,0,.35,1]},{'text':'正文','box':[.4,0,.6,1]}]}
    result,_=refine_layout_text(rgb,[item])
    assert len(result)==2
    assert result[0]['bold'] is True
    assert result[0]['font_size']==256


def test_frame_stroke_width_is_measured_from_pixels(tmp_path):
    from image2editable.layout_first import prepare_layout

    rgb=np.full((600,1200,3),255,np.uint8)
    cv2.rectangle(rgb,(100,100),(1100,500),(250,240,200),-1)
    cv2.rectangle(rgb,(100,100),(1100,500),(220,160,20),4)
    path=tmp_path/'frame.png';Image.fromarray(rgb).save(path)
    slide,_=prepare_layout(path,tmp_path/'assets',text_items=[])
    cards=[e for e in slide['visual_elements'] if e.get('role')=='card']
    assert cards
    assert 3<=cards[0]['shape']['line_width']<=6


def test_audit_does_not_call_foreign_icon_a_duplicate_card(tmp_path):
    from image2editable.layout_first import _audit_invariants

    rgba = np.zeros((100,160,4),np.uint8)
    rgba[20:80,110:150] = [20,40,90,255]
    path=tmp_path/'icon.png';Image.fromarray(rgba).save(path)
    elements=[{'object_id':'card','route':'native_shape','role':'card','bbox':[0,0,100,100],
               'shape':{'fill_rgb':[245,245,245]}},
              {'object_id':'icon','route':'raster_component','role':'visual',
               'component':{'path':str(path),'x':0,'y':0,'w':160,'h':100}}]
    assert _audit_invariants(elements,[],(200,200))['i3']==0


def test_three_stamp_series_preserves_filled_ink_and_route(tmp_path):
    from image2editable.layout_first import prepare_layout

    rgb = np.full((1000, 1800, 3), 255, np.uint8)
    items = []
    for n, y in enumerate((60, 370, 680)):
        cv2.rectangle(rgb, (100, y), (1700, y+260), (245, 248, 252), -1)
        points = np.array([[1350, y+65], [1610, y+50], [1630, y+190], [1370, y+210]])
        cv2.fillConvexPoly(rgb, points, (210, 65, 50))
        cv2.putText(rgb, 'ABC', (1390, y+155), cv2.FONT_HERSHEY_SIMPLEX, 1.5, (255,255,255), 4)
        items.append({'text': 'ABC', 'box': [1380,y+105,200,55], 'color':'#ffffff',
                      'font_size':32,'confidence':.99})
    path = tmp_path/'stamps.png'
    Image.fromarray(rgb).save(path)
    slide, _ = prepare_layout(path, tmp_path/'assets', text_items=items)
    stamps = [e for e in slide['visual_elements'] if e['object_id'].startswith('emblem_')]
    assert len(stamps) == 3
    for e in stamps:
        rgba = np.asarray(Image.open(e['component']['path']).convert('RGBA'))
        red = (rgba[:,:,0]>180)&(rgba[:,:,1]<100)&(rgba[:,:,3]>250)
        assert red.sum() > 20000


def test_duplicate_emblems_on_same_ink_cluster_collapse_to_one(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1000, 3), 255, np.uint8)
    blob = np.array([[(300, 280), (360, 80), (520, 120), (540, 420), (400, 540)]])
    cv2.fillConvexPoly(pixels, blob, (200, 120, 40))
    cv2.circle(pixels, (420, 300), 110, (60, 160, 240), -1)
    items = [{'text': '诊断', 'box': [380, 380, 140, 70], 'color': '#ffffff', 'font_size': 40, 'confidence': .95},
             {'text': '03', 'box': [330, 130, 120, 90], 'color': '#ffffff', 'font_size': 50, 'confidence': .95}]
    path = tmp_path / 'emblem.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items, card_mode='native')
    emblems = [e for e in slide['visual_elements'] if 'emblem' in e['object_id']]
    assert len(emblems) == 1


def test_sibling_card_series_suppresses_read_digit_when_others_unread(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1400, 3), 255, np.uint8)
    for x0 in (60, 730):
        cv2.rectangle(pixels, (x0, 80), (x0 + 560, 560), (252, 248, 240), -1)
    # card A (left): giant art digits stay unreadable; card B (right): same slot
    # got OCR'd -> mixed series must collapse the read item into a PNG too
    for x0 in (180, 330):
        cv2.rectangle(pixels, (x0, 140), (x0 + 110, 340), (143, 191, 249), -1)
    for x0 in (860, 1010):
        cv2.rectangle(pixels, (x0, 140), (x0 + 110, 340), (143, 191, 249), -1)
    for x0 in (60, 730):
        cv2.rectangle(pixels, (x0 + 120, 430), (x0 + 440, 500), (31, 95, 191), -1)
    items = [{'text': '03', 'box': [858, 138, 264, 204], 'color': '#8fbff9',
              'font_size': 80, 'confidence': .95},
             {'text': '标题一', 'box': [180, 430, 320, 70], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95},
             {'text': '标题二', 'box': [850, 430, 320, 70], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95}]
    path = tmp_path / 'series.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items, card_mode='native')
    texts = [i['text'] for i in slide['text_items']]
    assert '03' not in texts
    assert '标题一' in texts and '标题二' in texts
    assert audit['text_audit'].get('unit_suppressed', 0) >= 1


def test_inverted_highlight_card_keeps_series_membership(tmp_path):
    # a highlight column inverts BOTH card surface and digit color (light
    # digits on a dark card while siblings are dark on light). Series
    # membership is positional and scale-based only: any absolute element-to
    # -element color comparison ejects the recolored sibling and leaves the
    # read digit native while its twins ship as PNGs.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1400, 3), 255, np.uint8)
    cv2.rectangle(pixels, (60, 80), (620, 560), (250, 246, 240), -1)    # light card
    cv2.rectangle(pixels, (730, 80), (1290, 560), (28, 52, 150), -1)    # dark card
    for x0 in (190, 340):                                              # dark digits, OCR'd
        cv2.rectangle(pixels, (x0, 140), (x0 + 90, 330), (40, 70, 170), -1)
    for x0 in (870, 1020):                                             # light digits, unread
        cv2.rectangle(pixels, (x0, 140), (x0 + 90, 330), (235, 238, 255), -1)
    items = [{'text': '01', 'box': [188, 138, 254, 194], 'color': '#2846aa',
              'font_size': 80, 'confidence': .95}]
    path = tmp_path / 'inverted.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items,
                                  card_mode='native')
    texts = [i['text'] for i in slide['text_items']]
    assert '01' not in texts, \
        'read digit on the light card must follow its inverted sibling to PNG'
    assert audit['text_audit'].get('unit_suppressed', 0) >= 1


def test_free_sliver_near_owned_artwork_is_merged(tmp_path):
    # card-interior artwork lands in the owned residual pass; a sliver just
    # outside the card edge lands in the free pass. Same-pass merging leaves
    # the sliver as its own tiny object; the merge pool must span both passes
    # so the sliver attaches to the artwork (png2 p02 regression).
    from image2editable.layout_first import prepare_layout

    rng = np.random.default_rng(7)
    pixels = np.full((800, 1200, 3), 255, np.uint8)
    cv2.rectangle(pixels, (700, 200), (1100, 700), (30, 60, 160), -1)
    pixels[300:500, 730:930] = rng.integers(60, 200, (200, 200, 3), np.uint8)
    cv2.rectangle(pixels, (665, 400), (690, 440), (30, 30, 30), -1)
    path = tmp_path / 'sliver.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[],
                                  card_mode='native')
    visuals = [e for e in slide['visual_elements']
               if e.get('role') == 'visual' and 'component' in e]
    tiny = [e for e in visuals
            if max(e['component']['w'], e['component']['h']) < 50
            or e['component']['w'] * e['component']['h'] < 800 * 1200 * .0003]
    assert visuals, 'artwork should survive as a raster object'
    assert not tiny, f'independent tiny objects survived: {tiny}'


def test_visible_owned_sliver_without_merge_target_is_emitted(tmp_path):
    # a divider dash floating mid-card is an owned residual sliver with no
    # merge target nearby. Absorbing it into the flat card fill deletes the
    # pixels; only a sliver whose contrast is below the visible threshold may
    # be dropped (png2 p03 divider-dash regression).
    from image2editable.layout_first import prepare_layout

    rng = np.random.default_rng(7)
    pixels = np.full((1000, 1400, 3), 255, np.uint8)
    cv2.rectangle(pixels, (100, 100), (800, 800), (30, 60, 160), -1)
    pixels[200:400, 300:500] = rng.integers(60, 200, (200, 200, 3), np.uint8)
    cv2.rectangle(pixels, (300, 600), (360, 603), (240, 245, 255), -1)
    path = tmp_path / 'dash.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[],
                                  card_mode='native')
    visuals = [e for e in slide['visual_elements']
               if e.get('role') == 'visual' and 'component' in e]
    dash = [e for e in visuals
            if e['component']['y'] <= 620 and e['component']['y'] + e['component']['h'] >= 590
            and e['component']['w'] >= 40]
    assert dash, 'visible in-card divider dash was silently dropped'


def test_dark_card_with_painted_light_content_is_still_a_card(tmp_path):
    # white digits + icon painted on a dark card look like letterform counters to
    # the enclosed-hole check; a container this large must survive it, or the
    # emblem pass swallows the whole column into one flat PNG (png2 directory
    # page regression).
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1000, 3), 255, np.uint8)
    cv2.rectangle(pixels, (100, 60), (500, 540), (30, 70, 190), -1)
    cv2.rectangle(pixels, (160, 100), (260, 260), (255, 255, 255), -1)
    cv2.rectangle(pixels, (300, 100), (400, 260), (255, 255, 255), -1)
    cv2.circle(pixels, (300, 400), 60, (240, 245, 255), -1)
    path = tmp_path / 'dark.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], card_mode='native')
    cards = [o for o in audit['objects'] if o['role'] == 'card']
    assert any(c['bbox'][0] <= 120 and c['bbox'][2] >= 480 for c in cards)


def test_decoration_inside_text_box_on_card_is_released_as_object(tmp_path):
    # a badge frame painted through a native text item's box is not glyph ink;
    # masking the whole item box as text deletes it. The orange underline must
    # surface as its own residual object (png2 "本课重点" frame regression).
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1000, 3), 255, np.uint8)
    cv2.rectangle(pixels, (100, 60), (500, 540), (30, 70, 190), -1)
    cv2.putText(pixels, 'KEY', (180, 292), cv2.FONT_HERSHEY_SIMPLEX, 3.2,
                (255, 255, 255), 14, cv2.LINE_AA)
    cv2.rectangle(pixels, (170, 302), (430, 309), (252, 168, 16), -1)
    items = [{'text': 'KEY', 'box': [170, 220, 260, 90], 'color': '#ffffff',
              'font_size': 60, 'confidence': .95}]
    path = tmp_path / 'frame.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items, card_mode='native')
    texts = [i['text'] for i in slide['text_items']]
    assert 'KEY' in texts
    cards = [o for o in audit['objects'] if o['role'] == 'card']
    assert cards
    visuals = [e for e in slide['visual_elements'] if e.get('role') == 'visual']
    assert any(abs(v['component']['y'] - 302) < 30 for v in visuals), \
        'badge underline should ship as a raster object, not vanish into the text mask'


def test_unit_png_drops_foreign_color_slivers(tmp_path):
    # a suppressed digit unit clips its alpha to pixels matching the source
    # item color; a decoration sliver inside the OCR box must not ride along.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1400, 3), 255, np.uint8)
    for x0 in (60, 730):
        cv2.rectangle(pixels, (x0, 80), (x0 + 560, 560), (252, 248, 240), -1)
    for x0 in (180, 330):
        cv2.rectangle(pixels, (x0, 140), (x0 + 110, 340), (143, 191, 249), -1)
    for x0 in (860, 1010):
        cv2.rectangle(pixels, (x0, 140), (x0 + 110, 340), (143, 191, 249), -1)
    for x0 in (60, 730):
        cv2.rectangle(pixels, (x0 + 120, 430), (x0 + 440, 500), (31, 95, 191), -1)
    cv2.rectangle(pixels, (1100, 145), (1200, 160), (252, 168, 16), -1)
    items = [{'text': '03', 'box': [858, 138, 350, 204], 'color': '#8fbff9',
              'font_size': 80, 'confidence': .95},
             {'text': '标题一', 'box': [180, 430, 320, 70], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95},
             {'text': '标题二', 'box': [850, 430, 320, 70], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95}]
    path = tmp_path / 'sliver.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=items, card_mode='native')
    units = [e for e in slide['visual_elements'] if 'unit' in e['object_id']]
    assert units
    for unit in units:
        im = np.asarray(Image.open(unit['component']['path']))
        a = im[..., 3]
        orange = (im[..., 0] > 200) & (im[..., 1] > 80) & (im[..., 1] < 200) & (im[..., 2] < 90) & (a > 0)
        # a few boundary pixels survive where the sliver touches a glyph edge;
        # the sliver itself must be gone
        assert orange.sum() < 250, f"{unit['object_id']} carries foreign-color pixels"


def test_row_chain_box_ignores_row_below():
    # digit seeds must not union a wide icon that only overlaps them in x --
    # the unit PNG ("01") stays one row, the icon ships as its own object.
    from image2editable.layout_first import _row_chain_box

    comps = [{'bbox': (100, 100, 160, 240)}, {'bbox': (185, 120, 250, 240)},
             {'bbox': (90, 300, 350, 480)}]
    box, members = _row_chain_box(comps, [(100, 100, 160, 240)])
    assert box == (100, 100, 250, 240)
    assert [m['bbox'] for m in members] == [(100, 100, 160, 240), (185, 120, 250, 240)]


def test_source_backdrop_mode_extracts_white_cards_and_removes_their_old_pixels(tmp_path):
    from image2editable.layout_first import prepare_layout
    from image2editable.shape_analysis import _rounded_rectangle

    pixels = np.full((360, 640, 3), [20, 100, 220], np.uint8)
    pixels[:, :200] = np.random.default_rng(4).integers(50, 200, (360, 200, 3), np.uint8)
    pill = _rounded_rectangle(70, 220, 34)
    pixels[180:250, 330:550][pill] = 255
    path = tmp_path / 'cover.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[], background_mode='source', card_mode='png')
    assert audit['background_kind'] in {'source_cleaned', 'source_fill'}
    assert audit['card_images'] == 1
    assert slide.get('background_rgb') == [20, 100, 220] or (
        'background_path' in slide and np.max(np.abs(np.asarray(Image.open(slide['background_path']))[210, 400].astype(int)-[20, 100, 220])) <= 3)
    pictures = [e for e in slide['visual_elements'] if e['route'] == 'raster_component']
    assert pictures
    assert not any(e['component']['w'] >= 640 and e['component']['h'] >= 360 for e in pictures)


def test_source_backdrop_splits_white_cards_joined_by_page_gutter(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((720, 1280, 3), (247, 250, 254), np.uint8)
    pixels[:, :280] = np.random.default_rng(2).integers(40, 190, (720, 280, 3), np.uint8)
    pixels[:70] = (12, 70, 180)
    cards = ((90, 230), (280, 420), (470, 610))
    texts = []
    for index, (top, bottom) in enumerate(cards, 1):
        pixels[top:bottom, 360:1240] = 255
        pixels[top + 20:bottom - 20, 370:520] = (40, 110, 210)
        cv2.rectangle(pixels, (80, top + 40), (250, top + 90), (255, 255, 255), -1)
        texts.append({'box': [540, top + 50, 400, 36], 'text': f'BODY {index}', 'font_size': 18,
                      'font': 'Arial', 'color': '#203040', 'bold': False})
        texts.append({'box': [390, top + 50, 110, 32], 'text': f'L{index}', 'font_size': 16,
                      'font': 'Arial', 'color': '#ffffff', 'bold': True})
    pixels[248:254, 360:1240] = (192, 215, 248)
    pixels[448:454, 360:1240] = (192, 215, 248)
    path = tmp_path / 'gutter-cards.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=texts, background_mode='source')
    cards = [e for e in slide['visual_elements']
             if e.get('role') == 'card' and e['bbox'][2] - e['bbox'][0] > 400]
    assert len(cards) >= 3
    tops = sorted(e['bbox'][1] for e in cards[:3])
    assert tops[1] - tops[0] > 80 and tops[2] - tops[1] > 80
    grouped = {item for members in slide['object_groups'].values() for item in members}
    assert {'text_0001', 'text_0003', 'text_0005'} <= grouped
    assert 'background_rgb' in slide
    assert np.max(np.abs(np.array(slide['background_rgb']) - [247, 250, 254])) <= 8
    pictures = [e for e in slide['visual_elements'] if e['route'] == 'raster_component']
    assert pictures
    assert all(e['component']['w'] < 1280 or e['component']['h'] < 720 for e in pictures)
    assert any(e['component']['x'] < 280 and e['component']['h'] > 200 for e in pictures)
    assert 'background_path' not in slide


def test_source_backdrop_mode_does_not_put_flat_page_cards_back_into_background(tmp_path):

    from image2editable.layout_first import prepare_layout

    source, text = _source(tmp_path)
    slide, audit = prepare_layout(source, tmp_path / 'assets', text_items=text, background_mode='source')
    assert audit['background_kind'] == 'flat'
    assert audit['native_shapes'] == 2
    assert slide['background_rgb'] == [255, 255, 255]


def test_repeated_chevron_decoration_is_one_complete_transparent_object(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((720, 1280, 3), 255, np.uint8)
    for x, color in zip((20, 120, 220), ([0, 102, 221], [76, 142, 239], [163, 195, 245])):
        points = np.array([[x, 80], [x+30, 80], [x+70, 120], [x+30, 160], [x, 160], [x+40, 120]])
        cv2.fillPoly(pixels, [points], color)
    path = tmp_path / 'chevrons.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    pictures = [e for e in slide['visual_elements'] if e['route'] == 'raster_component']
    assert len(pictures) == 1
    assert audit['visual_parts_merged'] == 3
    c = pictures[0]['component']
    with Image.open(c['path']) as image:
        rgba = np.asarray(image)
    assert rgba[120-c['y'], 105-c['x'], 3] == 0


def test_different_small_icons_are_not_merged_by_proximity(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((720, 1280, 3), 255, np.uint8)
    cv2.fillPoly(pixels, [np.array([[20, 80], [50, 80], [90, 120], [50, 160], [20, 160], [60, 120]])], [20, 90, 210])
    cv2.fillPoly(pixels, [np.array([[130, 80], [170, 160], [100, 160]])], [20, 90, 210])
    path = tmp_path / 'different.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert len([e for e in slide['visual_elements'] if e['route'] == 'raster_component']) == 2


def test_interleaved_chevrons_merge_despite_overlapping_bounding_boxes(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((720, 1280, 3), 255, np.uint8)
    for x in (20, 70, 120):
        points = np.array([[x, 80], [x+30, 80], [x+70, 120], [x+30, 160], [x, 160], [x+40, 120]])
        cv2.fillPoly(pixels, [points], [20, 90, 210])
    path = tmp_path / 'interleaved.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[])
    assert audit['visual_parts_merged'] == 3
    assert len([e for e in slide['visual_elements'] if e['route'] == 'raster_component']) == 1


def test_wide_outlined_card_is_not_an_endpointless_native_line(tmp_path):
    from image2editable.layout_first import convert_layout, prepare_layout

    fixture = Path(__file__).resolve().parent / 'fixtures' / 'wide-outlined-card.png'
    path = tmp_path / 'wide-card.png'
    path.write_bytes(fixture.read_bytes())
    text = [{'box': [80, 80, 180, 40], 'text': 'CASE', 'font_size': 18,
             'font': 'Arial', 'color': '#ffffff', 'bold': True}]
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=text, background_mode='source')
    lines = [e for e in slide['visual_elements']
             if e.get('route') == 'native_shape' and e.get('shape', {}).get('shape_type') == 'line']
    assert all('line_start' in e['shape'] and 'line_end' in e['shape'] for e in lines)
    cards = [e for e in slide['visual_elements'] if e.get('role') == 'card' and e.get('route') == 'native_shape']
    assert cards
    assert all(e['shape']['shape_type'] != 'line' for e in cards)
    cache = tmp_path / 'text.json'
    cache.write_text(json.dumps({'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                                 'text_items': text}), encoding='utf-8')
    convert_layout([path], output_path=tmp_path / 'wide-card.pptx', run_dir=tmp_path / 'run',
                   text_cache=cache, slide_size='16:9', background_mode='source', render=False)


def test_enclosed_text_item_is_dropped(tmp_path):
    from image2editable.layout_first import refine_layout_text

    rgb = np.full((600, 1200, 3), 255, np.uint8)
    items = [
        {'box': [100, 200, 30, 60], 'text': '引', 'font_size': 20, 'font': 'Arial',
         'color': '#0039b6', 'bold': True, 'confidence': 0.9},
        {'box': [100, 200, 240, 62], 'text': '张女士，', 'font_size': 20, 'font': 'Arial',
         'color': '#222222', 'bold': False, 'confidence': 0.95},
    ]
    refined, audit = refine_layout_text(rgb, items)
    assert [i['text'] for i in refined] == ['张女士，']
    assert audit['overlap_dropped'] == 1


def test_card_edge_fringe_line_is_not_emitted(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((500, 800, 3), 255, np.uint8)
    pixels[120:340, 100:700] = (235, 243, 251)
    pixels[115:117, 100:700] = (236, 244, 252)
    pixels[117:119, 100:700] = (210, 226, 246)
    path = tmp_path / 'fringe.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    lines = [e for e in slide['visual_elements']
             if e.get('route') == 'native_shape' and e.get('shape', {}).get('shape_type') == 'line']
    assert not lines
    assert any(e['route'] == 'raster_component'
               and e['component']['y'] <= 117 < e['component']['y'] + e['component']['h']
               for e in slide['visual_elements'])


def test_sandwiched_two_tone_separator_becomes_raster(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((700, 1200, 3), 255, np.uint8)
    pixels[100:280, 200:1000] = (235, 243, 251)
    pixels[320:500, 200:1000] = (235, 243, 251)
    pixels[290:292, 200:1000] = (192, 215, 248)
    pixels[292:294, 200:1000] = (121, 170, 241)
    pixels[294:296, 200:1000] = (192, 215, 248)
    path = tmp_path / 'sandwich.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    lines = [e for e in slide['visual_elements']
             if e.get('route') == 'native_shape' and e.get('shape', {}).get('shape_type') == 'line']
    assert not lines
    assert any(e['route'] == 'raster_component'
               and e['component']['x'] <= 210 and e['component']['x'] + e['component']['w'] >= 990
               and e['component']['y'] <= 292 < e['component']['y'] + e['component']['h']
               for e in slide['visual_elements'])


def test_full_width_band_emits_one_strip_not_lines(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1200, 3), 255, np.uint8)
    pixels[200:202] = (244, 248, 254)
    pixels[202:210] = (46, 125, 235)
    pixels[210:212] = (244, 248, 254)
    pixels[208:280, 500:700] = (46, 125, 235)
    pixels[320:520, 100:900] = (235, 243, 251)
    path = tmp_path / 'band.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=[])
    lines = [e for e in slide['visual_elements']
             if e.get('route') == 'native_shape' and e.get('shape', {}).get('shape_type') == 'line']
    assert not lines
    strips = [e for e in slide['visual_elements'] if e['object_id'].startswith('band_')]
    assert len(strips) == 1
    assert strips[0]['bbox'][0] <= 4 and strips[0]['bbox'][2] >= 1196
    assert strips[0]['bbox'][1] <= 202 and strips[0]['bbox'][3] >= 210


def test_sibling_card_labels_unify_typography(tmp_path):
    from image2editable.layout_first import prepare_layout

    pixels = np.full((900, 1400, 3), 255, np.uint8)
    for top in (80, 330, 580):
        pixels[top:top+200, 200:1200] = (240, 246, 252)
        cv2.putText(pixels, 'AB', (245, top + 85), cv2.FONT_HERSHEY_SIMPLEX, 1.6,
                    (0, 57, 182), 3, cv2.LINE_AA)
    text = [
        {'box': [240, 120, 190, 55], 'text': '病例：', 'font_size': 20, 'font': 'Arial',
         'color': '#0039b6', 'bold': True, 'confidence': .9},
        {'box': [240, 370, 190, 55], 'text': '主诉：', 'font_size': 19.5, 'font': 'Arial',
         'color': '#0039b6', 'bold': True, 'confidence': .9},
        {'box': [240, 620, 190, 55], 'text': '经过：', 'font_size': 19, 'font': 'Arial',
         'color': '#003ab6', 'bold': False, 'confidence': .9},
    ]
    path = tmp_path / 'series.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=text)
    kept = [i for i in slide['text_items'] if i['text'] and i['box'][1] > 60]
    assert kept
    assert {i['bold'] for i in kept} == {True}
    assert len({i['font_size'] for i in kept}) == 1
    assert len({i['color'] for i in kept}) == 1


def test_long_sentence_with_unread_glyph_stays_native(tmp_path):
    # p04 B-any-01c regression: an OCR-readable body line sat beside one
    # unread same-color glyph (quote/punctuation) and the whole sentence was
    # rasterized, violating text rule 1. Partial-unit suppression only covers
    # compact text units ("70%" style); a line many glyph-heights wide is a
    # sentence and must stay native.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((600, 1600, 3), 255, np.uint8)
    cv2.rectangle(pixels, (80, 120), (1520, 460), (252, 248, 240), -1)
    cv2.rectangle(pixels, (140, 200), (1420, 260), (31, 95, 191), -1)
    cv2.rectangle(pixels, (1450, 210), (1480, 240), (31, 95, 191), -1)
    items = [{'text': '面对这位接触性出血的育龄女性患者你会想到什么',
              'box': [140, 200, 1280, 60], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95}]
    path = tmp_path / 'sentence.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=items,
                              card_mode='native')
    texts = [i['text'] for i in slide['text_items']]
    assert any('接触性出血' in text for text in texts), \
        'a wide OCR line must stay native text, not a raster unit'


def test_wide_body_line_is_not_dragged_into_label_series(tmp_path):
    # p02 B-any-01c regression: the slot-member match accepted a body line as
    # the series twin of a compact label, so label suppression rasterized a
    # full sentence. Slot members must be compact text units as well.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((1100, 2200, 3), 255, np.uint8)
    cv2.rectangle(pixels, (560, 80), (2060, 480), (252, 248, 240), -1)
    cv2.rectangle(pixels, (60, 140), (480, 420), (232, 224, 208), -1)
    cv2.rectangle(pixels, (1600, 190), (1800, 300), (31, 95, 191), 8)
    cv2.rectangle(pixels, (180, 220), (460, 290), (31, 95, 191), -1)
    cv2.rectangle(pixels, (1000, 200), (1900, 260), (31, 95, 191), -1)
    items = [{'text': '标签', 'box': [180, 220, 280, 70], 'color': '#1f5fbf',
              'font_size': 45, 'confidence': .95},
             {'text': '这是一行正文不应该被拖进标签的系列栅格化',
              'box': [1000, 200, 900, 60], 'color': '#1f5fbf',
              'font_size': 40, 'confidence': .95}]
    path = tmp_path / 'series.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=items,
                              card_mode='native')
    texts = [i['text'] for i in slide['text_items']]
    assert any('不应该被拖进' in text for text in texts), \
        'a wide body line must not be suppressed as a label-series member'


def test_card_frame_chain_never_swallows_its_host_card(tmp_path):
    # p04 B-any-01c regression: a card's own frame ink comp seeded an unread
    # row chain that covered the whole card -> the card shipped twice (native
    # shape + whole-card PNG, I3 violation). A unit may never swallow a card.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((900, 2400, 3), 255, np.uint8)
    for x0 in (80, 1300):
        cv2.rectangle(pixels, (x0, 200), (x0 + 1020, 700), (255, 244, 191), -1)
        cv2.rectangle(pixels, (x0 + 8, 208), (x0 + 1012, 692), (235, 205, 60), 12)
    for x0 in (400, 1620):
        cv2.rectangle(pixels, (x0, 380), (x0 + 240, 540), (90, 90, 90), -1)
    path = tmp_path / 'framed.png'
    Image.fromarray(pixels).save(path)
    slide, audit = prepare_layout(path, tmp_path / 'assets', text_items=[],
                                  card_mode='native')
    cards = [e for e in slide['visual_elements'] if e.get('role') == 'card']
    assert cards, 'the yellow cards should be detected'
    assert all(8 <= c['shape']['line_width'] <= 16 for c in cards if c['bbox'][2]-c['bbox'][0] > 800)
    for e in slide['visual_elements']:
        if 'component' not in e:
            continue
        comp = e['component']
        ex0, ey0 = comp['x'], comp['y']
        ex1, ey1 = ex0 + comp['w'], ey0 + comp['h']
        for card in cards:
            cx0, cy0, cx1, cy1 = card['bbox']
            inter = (max(0, min(ex1, cx1) - max(ex0, cx0))
                     * max(0, min(ey1, cy1) - max(ey0, cy0)))
            assert inter < .6 * (cx1 - cx0) * (cy1 - cy0), \
                f'{e["object_id"]} raster-covers card {card["object_id"]}'


def test_row_chain_box_stops_at_element_gaps():
    # a row chain is one element: growth is bounded by member height and the
    # chain box may never exceed a few glyph heights, so an icon or stamp
    # sitting element-gaps away is never welded into the row (p02 regression).
    from image2editable.layout_first import _row_chain_box

    # element-ladder: each next comp sits just inside the growing box reach,
    # so an unbounded chain sweeps the whole row; element-bounded growth must
    # stop before crossing gaps larger than the member height.
    comps = [{'bbox': (x, 100, x + w_, 220), 'area': 3000, 'h': 120,
              'color': np.array([30, 60, 160]), 'mask': np.ones((120, 60), bool)}
             for x, w_ in ((100, 60), (190, 90), (330, 90), (500, 120),
                           (700, 120), (1100, 120), (1500, 120))]
    box, members = _row_chain_box(comps, [(100, 100, 160, 220)])
    assert box[2] - 100 <= 120 * 5, \
        f'chain welded a whole row of separate elements: {box}'


def test_three_pill_rows_at_page_bottom_do_not_seed_outer_units(tmp_path):
    from image2editable.layout_first import _suppress_mixed_unit_text

    rgb = np.full((720, 1200, 3), 255, np.uint8)
    regions = []
    for y in (20, 260, 500):
        parent_mask = np.zeros((720, 1200), np.uint8)
        parent_mask[y:y+220, 201:1180] = 255
        regions.append({'bbox': [201, y, 1180, y+220], 'transform': None,
                        'aligned_mask': parent_mask[y:y+220, 201:1180],
                        'surface': (rgb[y:y+220, 201:1180].copy(),),
                        'shape': {'shape_type': 'rectangle', 'fill_rgb': [255]*3}})
        x = 200 if y == 500 else 202
        mask = np.zeros((720, 1200), np.uint8)
        cv2.rectangle(mask, (x+35, y+40), (x+245, y+180), 255, -1)
        cv2.circle(mask, (x+35, y+110), 70, 255, -1)
        cv2.circle(mask, (x+245, y+110), 70, 255, -1)
        mask[:, :x] = 0
        rgb[mask != 0] = [47, 107, 208]
        regions.append({'bbox': [x, y+40, x+315, y+181], 'transform': None,
                        'aligned_mask': mask[y+40:y+181, x:x+315],
                        'surface': (np.full((141, 315, 3), [47, 107, 208], np.uint8),),
                        'shape': {'shape_type': 'rounded_rectangle', 'fill_rgb': [47, 107, 208]}})
        cv2.rectangle(rgb, (540, y+70), (650, y+160), (30, 60, 130), -1)
    _, _, specs, _ = _suppress_mixed_unit_text(rgb, regions, [], np.zeros((720, 1200), np.uint8))
    assert not any(s['box'][0] < 515 and s['box'][2] > 540 for s in specs)


def test_disconnected_icon_on_pale_panel_is_one_faithful_png(tmp_path):
    from image2editable.layout_first import prepare_layout

    rgb = np.full((800, 1400, 3), 255, np.uint8)
    for y in (100, 430):
        cv2.rectangle(rgb, (100, y), (1300, y+300), (249, 249, 249), -1)
        cv2.rectangle(rgb, (450, y+60), (590, y+240), (223, 237, 251), -1)
        cv2.line(rgb, (480, y+80), (480, y+110), (155, 170, 191), 8)
        cv2.ellipse(rgb, (520, y+150), (35, 40), 0, 0, 180, (35, 55, 80), 8)
        cv2.circle(rgb, (555, y+215), 12, (35, 55, 80), 5)
    source = tmp_path / 'panels.png'
    Image.fromarray(rgb).save(source)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=[])
    for y in (100, 430):
        parts = []
        for e in slide['visual_elements']:
            c = e.get('component')
            if c and c['x'] < 590 and c['x']+c['w'] > 450 and c['y'] < y+240 and c['y']+c['h'] > y+60:
                parts.append(c)
        assert len(parts) == 1, parts
        c = parts[0]
        with Image.open(c['path']) as image:
            rgba = np.asarray(image.convert('RGBA'))
        for x, yy in ((480, y+90), (520, y+190), (555, y+203)):
            assert rgba[yy-c['y'], x-c['x'], 3] == 255
            assert np.array_equal(rgba[yy-c['y'], x-c['x'], :3], rgb[yy, x])


def test_nested_card_surface_never_seeds_outer_unit(tmp_path):
    # p02 B-any-01d regression: a nested label pill's surface is ink vs the
    # outer row card, so the whole pill became one giant comp and
    # _row_chain_box welded the pill + the neighbouring icon + a stamp into a
    # row-wide unit PNG. A nested card renders itself; its pixels must not
    # seed the outer region's unread-ink comps.
    from image2editable.layout_first import prepare_layout

    pixels = np.full((1000, 2400, 3), 255, np.uint8)
    for y0 in (100, 500):
        # sibling row card
        cv2.rectangle(pixels, (200, y0), (2200, y0 + 300), (252, 248, 240), -1)
        # nested label pill on the left of each row
        cv2.rectangle(pixels, (300, y0 + 60), (740, y0 + 240), (235, 240, 250), -1)
        cv2.rectangle(pixels, (300, y0 + 60), (740, y0 + 240), (60, 90, 180), 6)
        # an icon far to the right on the same row card
        cv2.circle(pixels, (1600, y0 + 150), 90, (40, 40, 40), -1)
        cv2.circle(pixels, (1600, y0 + 150), 50, (252, 248, 240), -1)
    # unread glyph marks inside the lower pill (a stamp OCR could not read)
    cv2.putText(pixels, 'XYZ', (330, 720), cv2.FONT_HERSHEY_SIMPLEX, 2.0,
                (60, 90, 180), 8, cv2.LINE_AA)
    items = [{'text': '懂理论', 'box': [330, 160, 380, 100], 'color': '#3c5ab4',
              'font_size': 60, 'confidence': .95}]
    path = tmp_path / 'nested.png'
    Image.fromarray(pixels).save(path)
    slide, _ = prepare_layout(path, tmp_path / 'assets', text_items=items,
                              card_mode='native')
    cards = [e for e in slide['visual_elements'] if e.get('role') == 'card']
    pills = [c for c in cards if c['bbox'][2] - c['bbox'][0] < 1000]
    assert len(pills) >= 2, f'expected two nested pill cards, got {cards}'
    for e in slide['visual_elements']:
        if 'component' not in e:
            continue
        comp = e['component']
        ex0, ey0 = comp['x'], comp['y']
        ex1, ey1 = ex0 + comp['w'], ey0 + comp['h']
        for pill in pills:
            px0, py0, px1, py1 = pill['bbox']
            inter = (max(0, min(ex1, px1) - max(ex0, px0))
                     * max(0, min(ey1, py1) - max(ey0, py0)))
            frac = inter / ((px1 - px0) * (py1 - py0))
            assert frac < .6, \
                f'{e["object_id"]} covers {frac:.0%} of nested card {pill["object_id"]}'


def test_unified_series_size_survives_assembler_width_fit():
    # the assembler shrinks each run's font until the text fits its OCR box;
    # a median-unified size that only fits SOME members renders different pt
    # per member (v27 p04: items all 19.2 but pptx showed 15.28/16.83/16.77).
    # The unified size must already sit under every member's own width clamp.
    from image2editable.layout_first import _unify_series_styles
    from scripts.ppt_assemble import _fit_ocr_font_size

    page_w = 3840
    regions = [{'bbox': (0, 0, 1900, 700), 'transform': None,
                'shape': {'fill_rgb': (252, 248, 240)}},
               {'bbox': (0, 760, 1900, 1400), 'transform': None,
                'shape': {'fill_rgb': (252, 248, 240)}}]
    items = [{'text': '就诊经过：', 'box': [200, 300, 250, 90],
              'color': '#0039b6', 'font': 'Microsoft YaHei',
              'font_size': 60, 'bold': True},
             {'text': '思', 'box': [200, 1040, 250, 90],
              'color': '#0039b6', 'font': 'Microsoft YaHei',
              'font_size': 40, 'bold': True}]
    _unify_series_styles(items, regions, page=(page_w, 2160))
    sizes = {round(it['font_size'], 3) for it in items}
    assert len(sizes) == 1
    for it in items:
        width_in = it['box'][2] * (40.0 / 3.0) / page_w
        fitted = _fit_ocr_font_size(it['text'], it.get('font', 'Microsoft YaHei'),
                                    it.get('bold', False), it.get('italic', False),
                                    it['font_size'], width_in)
        assert fitted == it['font_size'], \
            'the unified size is still per-item clamped -> renders uneven'


def test_unify_series_skips_sentence_wide_body_line():
    # p04 B-any-01d regression: '不好意思就医，目...' (width > 6 glyph
    # heights) shared the normalized slot of '课堂提问：' on a sibling card
    # and unioned into the blue label series -> recoloured #193778.
    # _unify_series_styles must keep the suppression pass's compactness gate:
    # a line many glyph-heights wide is a sentence, never a series member.
    from image2editable.layout_first import _unify_series_styles

    regions = [{'bbox': (100, 200, 900, 700), 'transform': None,
                'shape': {'fill_rgb': (252, 248, 240)}},
               {'bbox': (1000, 200, 1800, 700), 'transform': None,
                'shape': {'fill_rgb': (252, 248, 240)}}]
    items = [{'text': '课堂提问：', 'box': [150, 300, 250, 80],
              'color': '#0039b6', 'font': 'Microsoft YaHei',
              'font_size': 40, 'bold': True},
             {'text': '不好意思就医，目前已经两个月了', 'box': [1050, 300, 700, 60],
              'color': '#33383b', 'font': 'Microsoft YaHei',
              'font_size': 36, 'bold': False}]
    _unify_series_styles(items, regions, page=(2400, 1200))
    assert items[1]['color'] == '#33383b'
    assert items[1]['font_size'] == 36


def test_source_title_recovery_rejects_crop_sized_recognition_box(tmp_path, monkeypatch):
    from image2editable.layout_first import _recover_source_titles

    monkeypatch.setattr('scripts.art_text.estimate_art_text_runs', lambda *a, **k: [{'text': 'TITLE'}])
    rgb = np.full((700, 1400, 3), 240, np.uint8)
    cv2.rectangle(rgb, (600, 220), (900, 330), (20, 20, 20), -1)
    items = [{'text': 'header', 'box': [600, 50, 400, 70]},
             {'text': 'footer', 'box': [600, 330, 400, 70]}]
    def reader(path, **kwargs):
        w, h = Image.open(path).size
        return [{'text': 'TITLE', 'box': [0, 0, w, h], 'confidence': .99}], None
    result, _ = _recover_source_titles(rgb, items, tmp_path, reader)
    assert result == items


def test_refine_trims_overlapping_padding_without_deleting_words():
    from image2editable.layout_first import refine_layout_text

    items = [{'text': '张女士，', 'box': [20, 20, 220, 50], 'color': '#333333',
              'words': [{'text': '张女士，', 'box': [.1, 0, .7, 1]}]},
             {'text': '36岁，', 'box': [220, 20, 180, 50], 'color': '#333333',
              'words': [{'text': '36岁，', 'box': [.1, 0, .8, 1]}]}]
    result, _ = refine_layout_text(np.full((100, 500, 3), 255, np.uint8), items)
    assert [v['text'] for v in result] == ['张女士，', '36岁，']
    assert result[0]['box'][0]+result[0]['box'][2] <= result[1]['box'][0]


def test_sentence_with_terminal_punctuation_is_not_a_compact_series():
    from image2editable.layout_first import _unify_series_styles

    regions = [{'bbox': [0, 0, 600, 300], 'shape': {'fill_rgb': [255, 255, 255]}},
               {'bbox': [700, 0, 1300, 300], 'shape': {'fill_rgb': [255, 255, 255]}}]
    items = [{'text': '课堂提问：', 'box': [30, 30, 240, 80], 'font_size': 28,
              'color': '#0039b6', 'bold': True},
             {'text': '不好意思就医，', 'box': [730, 30, 460, 80], 'font_size': 24,
              'color': '#33383b', 'bold': False}]
    _unify_series_styles(items, regions, page=(1400, 400))
    assert items[1]['color'] == '#33383b'
    assert items[1]['font_size'] == 24


def test_geometry_removes_thin_protrusion_from_rounded_card():
    from image2editable.layout_first import _geometry
    from image2editable.shape_analysis import _rounded_rectangle

    mask = np.zeros((260, 700), np.uint8)
    mask[70:230, 40:640] = _rounded_rectangle(160, 600, 35).astype(np.uint8) * 255
    cv2.line(mask, (150, 30), (150, 85), 255, 5)
    rgb = np.full((260, 700, 3), 255, np.uint8)
    rgb[mask != 0] = [240, 220, 180]
    measured = _geometry(rgb, mask, np.zeros(mask.shape, np.uint8))
    assert measured is not None
    assert measured['shape']['shape_type'] == 'rounded_rectangle'
    assert measured['bbox'][1] >= 68


def test_refine_trims_misread_tail_overlapping_next_prefix():
    from image2editable.layout_first import refine_layout_text

    rgb = np.full((160, 500, 3), 255, np.uint8)
    items = [
        {'text': '就医，目', 'box': [100, 50, 150, 36], 'confidence': .92,
         'words': [{'text': '就', 'box': [0, 0, .25, 1]},
                   {'text': '医', 'box': [.25, 0, .25, 1]},
                   {'text': '，', 'box': [.5, 0, .1, 1]},
                   {'text': '目', 'box': [.8, 0, .2, 1]}]},
        {'text': '自购', 'box': [220, 50, 100, 36], 'confidence': .95,
         'words': [{'text': '自', 'box': [0, 0, .5, 1]},
                   {'text': '购', 'box': [.5, 0, .5, 1]}]},
    ]
    refined, audit = refine_layout_text(rgb, items)
    assert [item['text'] for item in refined] == ['就医，', '自购']
    assert audit['duplicate_spans_removed'] == 1


def test_refine_drops_different_ocr_text_for_same_glyph_geometry():
    # p04 B-any-01d regression: adjacent OCR spans both read the same glyph
    # geometry as different text. Text equality is unavailable here; geometric
    # containment must remove the smaller duplicate without deleting a real
    # neighbouring span.
    from image2editable.layout_first import refine_layout_text

    rgb = np.full((160, 500, 3), 255, np.uint8)
    items = [
        {'text': '不好意思就医，目', 'box': [100, 50, 220, 36],
         'font_size': 20, 'color': '#33383b', 'confidence': .92},
        {'text': '目', 'box': [290, 50, 30, 36],
         'font_size': 20, 'color': '#33383b', 'confidence': .80},
        {'text': '之后', 'box': [350, 50, 60, 36],
         'font_size': 20, 'color': '#33383b', 'confidence': .90},
    ]
    refined, _ = refine_layout_text(rgb, items)
    texts = [item['text'] for item in refined]
    assert texts.count('目') == 0
    assert '不好意思就医，目' in texts
    assert '之后' in texts


def test_card_raster_dedupe_preserves_unique_foreign_ink(tmp_path):
    # p02 B-any-01d regression: a unit covering a pill/card also contains a
    # book icon or emblem outside the native card fill. Bbox mutual coverage
    # alone deleted the unit and the unique visual disappeared from the PPTX.
    from image2editable.layout_first import _dedupe_card_rasters

    card_path = tmp_path / 'card.png'
    Image.new('RGBA', (200, 100), (235, 240, 250, 255)).save(card_path)
    card = {'object_id': 'card_0001', 'role': 'card', 'route': 'native_shape',
            'bbox': [0, 0, 200, 100],
            'shape': {'fill_rgb': [235, 240, 250]}}
    unit_rgba = np.full((100, 200, 4), (235, 240, 250, 255), np.uint8)
    unit_rgba[30:70, 135:175] = (38, 58, 150, 255)
    unit_path = tmp_path / 'unit.png'
    Image.fromarray(unit_rgba).save(unit_path)
    unit = {'object_id': 'unit_0001', 'role': 'visual',
            'route': 'raster_component',
            'component': {'path': str(unit_path), 'x': 0, 'y': 0,
                          'w': 200, 'h': 100}}
    kept, dropped = _dedupe_card_rasters([card, unit])
    assert dropped == 0
    assert unit in kept


def test_card_raster_dedupe_drops_pure_card_fill(tmp_path):
    from image2editable.layout_first import _dedupe_card_rasters

    card = {'object_id': 'card_0001', 'role': 'card', 'route': 'native_shape',
            'bbox': [0, 0, 200, 100],
            'shape': {'fill_rgb': [235, 240, 250]}}
    unit_path = tmp_path / 'unit.png'
    Image.new('RGBA', (200, 100), (235, 240, 250, 255)).save(unit_path)
    unit = {'object_id': 'unit_0001', 'role': 'visual',
            'route': 'raster_component',
            'component': {'path': str(unit_path), 'x': 0, 'y': 0,
                          'w': 200, 'h': 100}}
    kept, dropped = _dedupe_card_rasters([card, unit])
    assert dropped == 1
    assert unit not in kept


def test_mixed_color_line_split_keeps_glyph_with_own_color():
    # p04 '血': a yellow glyph followed by a white run was sampled as white
    # because a few brighter fringe pixels out-scored the populous ink in the
    # decile-core vote. The split then assigned the glyph to the white chunk,
    # its box released the foreign-coloured ink and the glyph shipped twice.
    from image2editable.layout_first import refine_layout_text

    pixels = np.tile(np.array([60, 100, 200], np.uint8), (140, 760, 1))
    text = '接触性出血的育龄女性'
    words = []
    for i, ch in enumerate(text):
        x0 = 40 + i * 60
        pixels[74:104, x0 + 12:x0 + 48] = (255, 217, 23) if i < 5 else (255, 255, 255)
        if i == 4:
            for k in range(6):
                pixels[76 + k * 4:80 + k * 4, x0 + 50:x0 + 58] = 255
        words.append({'text': ch, 'box': [(x0 - 20) / 720, .1, 60 / 720, .8]})
    item = {'box': [20, 60, 720, 60], 'text': text, 'color': '#ffffff',
            'font_size': 40, 'words': words}
    items, audit = refine_layout_text(pixels, [item])
    joined = ''.join(i['text'] for i in items)
    assert joined == text
    yellow = [i for i in items if i['color'].lower() != '#ffffff']
    assert len(yellow) == 1 and yellow[0]['text'] == '接触性出血'


def test_contained_residual_record_absorbs_into_host_element(tmp_path):
    # p02 stamp bar: a record fully inside a raster element's box is the same
    # visual — emitting it as its own PNG draws the zone twice (I3).
    from image2editable.layout_first import _emit_residual

    unit = np.zeros((100, 200, 4), np.uint8)
    unit[10:90, 20:180] = (200, 30, 30, 255)
    host = tmp_path / 'unit_0001.png'
    Image.fromarray(unit).save(host)
    elements = [{'object_id': 'unit_0001', 'route': 'raster_component', 'role': 'visual',
                 'component': {'path': str(host), 'x': 0, 'y': 0, 'w': 200, 'h': 100}}]
    bar = np.full((60, 30, 4), (210, 40, 40, 255), np.uint8)
    rec = {'rgba': bar, 'x': 150, 'y': 20, 'w': 30, 'h': 60, 'owned': True}
    diff = np.zeros((600, 800), np.float32)
    text = np.zeros((600, 800), np.uint8)
    emitted = _emit_residual([rec], elements, tmp_path, diff, (600, 800), text, [])
    assert emitted == 0
    merged = np.asarray(Image.open(elements[0]['component']['path']))
    assert merged[30:70, 155:175, 3].mean() == 255


def test_record_inside_bigger_record_merges_before_emit(tmp_path):
    # p04 photo strip: the owned slice of one visual sits inside the free
    # record's box; two PNGs covering the same zone violates I3.
    from image2editable.layout_first import _emit_residual

    big = np.full((100, 200, 4), (240, 200, 180, 255), np.uint8)
    small = np.full((60, 40, 4), (240, 210, 190, 255), np.uint8)
    records = [{'rgba': big, 'x': 0, 'y': 0, 'w': 200, 'h': 100, 'owned': False},
               {'rgba': small, 'x': 150, 'y': 30, 'w': 40, 'h': 60, 'owned': True}]
    diff = np.zeros((600, 800), np.float32)
    text = np.zeros((600, 800), np.uint8)
    elements = []
    emitted = _emit_residual(records, elements, tmp_path, diff, (600, 800), text, [])
    assert emitted == 1
    comp = elements[0]['component']
    merged = np.asarray(Image.open(comp['path']))
    assert merged.shape[:2] == (100, 200)
    assert merged[40:80, 155:190, 3].mean() == 255


def test_audit_i1_solid_underlay_not_flagged(tmp_path):
    # p05 chevron nav: inpainting erases glyph RGB but keeps alpha opaque, so a
    # solid-colour underlay legitimately covers the glyph zone.  I1 must only
    # fire when the PNG actually re-prints the glyph ink colour.
    from image2editable.layout_first import _audit_invariants

    src = np.full((120, 320, 3), (228, 236, 248), np.uint8)
    src[40:80, 60:260] = (90, 143, 226)  # dark blue glyph strokes
    underlay = np.full((120, 320, 4), (228, 236, 248, 255), np.uint8)
    path = tmp_path / 'object_0001.png'
    Image.fromarray(underlay).save(path)
    elements = [{'route': 'raster_component',
                 'component': {'path': str(path), 'x': 0, 'y': 0, 'w': 320, 'h': 120}}]
    items = [{'box': (30, 20, 260, 80)}]
    audit = _audit_invariants(elements, items, (120, 320), src)
    assert audit['i1'] == 0


def test_audit_i1_reprinted_glyph_flagged(tmp_path):
    # A residual PNG that carries the same ink as the source glyphs is a real
    # duplicate: same positions, same colour.
    from image2editable.layout_first import _audit_invariants

    src = np.full((120, 320, 3), (228, 236, 248), np.uint8)
    src[40:80, 60:260] = (90, 143, 226)
    dup = np.zeros((120, 320, 4), np.uint8)
    dup[40:80, 60:260] = (90, 143, 226, 255)  # re-printed glyph ink
    path = tmp_path / 'object_0002.png'
    Image.fromarray(dup).save(path)
    elements = [{'route': 'raster_component',
                 'component': {'path': str(path), 'x': 0, 'y': 0, 'w': 320, 'h': 120}}]
    items = [{'box': (30, 20, 260, 80)}]
    audit = _audit_invariants(elements, items, (120, 320), src)
    assert audit['i1'] == 1


def test_text_color_picks_dominant_ink_over_foreign_fringe():
    # A bright foreign sliver inside the crop (the neighbouring run's AA spill)
    # out-contrasts the real ink on a dark card; the dominant ink bucket must
    # win before the top-decile core vote.  Lives in the layout-first-private
    # sampler since the shared scripts module was restored to c162378.
    from image2editable.text_color import sample_text_color

    region = np.full((100, 60, 3), (60, 100, 200), np.uint8)
    region[10:92, 8:38] = (255, 217, 23)
    for k in range(5):
        region[12 + k * 16:20 + k * 16, 44:52] = 255
    color = sample_text_color(region)
    actual = tuple(int(color[i:i + 2], 16) for i in (1, 3, 5))
    assert actual == pytest.approx((255, 217, 23), abs=30)


# ---------------------------------------------------------------------------
# B-any-09: tail-glyph twins, pale-card fill recovery, residual compositing
# ---------------------------------------------------------------------------


def test_bany09_tail_punct_twin_is_erased_not_preserved():
    from image2editable.layout_first import _preserve_straddling_marks
    # run04 p04: the source '。' of '接触血液体液时佩戴。' straddles the frame's
    # right edge (ink x~3089-3108, box ends 3095).  It is the glyph the native
    # text redraws — not a neighbour decoration — so the whole component must
    # go to cleaning or it double-prints as '。。'.
    zone = np.full((60, 140, 3), 235, np.uint8)
    text = np.zeros((60, 140), np.uint8)
    text[8:52, 8:102] = 255                       # mask ≈ item box dilated +2
    cv2.circle(zone, (96, 30), 10, (40, 40, 120), 3)   # '。' ink x86..106
    ring = ((np.indices(zone.shape[:2])[1] - 96) ** 2
            + (np.indices(zone.shape[:2])[0] - 30) ** 2 <= 196)
    ring_comp = ring & (zone.max(axis=2) < 160)
    cleaning = text != 0
    items = [{'text': '……佩戴。', 'box': [0, 8, 100, 44]}]
    kept = _preserve_straddling_marks(
        cleaning.copy(), zone, text, items, (0, 0))
    # every ring pixel joins the clean set — text repaints it, not the PNG
    assert np.all(kept[ring_comp])
    # same geometry but the line ends in a hanzi: the ring is a neighbour
    # mark with no text to redraw it, so it survives whole as before
    kept = _preserve_straddling_marks(
        cleaning.copy(), zone, text,
        [{'text': '……佩戴', 'box': [0, 8, 100, 44]}], (0, 0))
    assert not np.any(kept & ring_comp)


def test_bany09_run04_p04_baked_tail_period_is_gone(tmp_path):
    from image2editable.layout_first import prepare_layout
    source, items = _bany06_fixture('run04', 4)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=items,
                              background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    texts = [it['text'] for it in slide['text_items']]
    assert '接触血液体液时佩戴。' in texts
    assert '喷溅风险操作时佩戴。' in texts
    # the zones hold only the source '。' rings, so any dark ink a PNG still
    # carries there is the baked glyph — it would re-print beside the live
    # text's own period and render as '。。'
    for element in slide['visual_elements']:
        c = element.get('component')
        if not c:
            continue
        for x0, y0, x1, y1 in ((3080, 1145, 3140, 1190),
                               (3085, 1845, 3140, 1915)):
            ix0, iy0 = max(x0, c['x']), max(y0, c['y'])
            ix1, iy1 = min(x1, c['x'] + c['w']), min(y1, c['y'] + c['h'])
            if ix1 <= ix0 or iy1 <= iy0:
                continue
            rgba = np.asarray(Image.open(c['path']).convert('RGBA'))
            zone = rgba[iy0 - c['y']:iy1 - c['y'], ix0 - c['x']:ix1 - c['x']]
            baked = (zone[:, :, :3].max(axis=2) < 200) & (zone[:, :, 3] > 100)
            assert not np.any(baked), \
                (element['object_id'], int(baked.sum()))


def test_bany09_run04_p01_no_page_sized_frame(tmp_path):
    from image2editable.layout_first import prepare_layout
    source, items = _bany06_fixture('run04', 1)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=items,
                              background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    for element in slide['visual_elements']:
        shape = element.get('shape') or {}
        stroke = shape.get('stroke_rgb')
        if stroke is not None:
            assert all(isinstance(v, int) and 0 <= v <= 255 for v in stroke), \
                (element['object_id'], stroke)
        if element.get('role') == 'frame':
            x0, y0, x1, y1 = element['bbox']
            assert not (x1 - x0 >= 3830 and y1 - y0 >= 2150), \
                element['object_id']


def test_bany09_flatbg_p03_pale_cards_emit_fills(tmp_path):
    from image2editable.layout_first import prepare_layout
    source, items = _bany06_fixture('flatbg', 3)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=items,
                              background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    fills = [e for e in slide['visual_elements']
             if e['route'] == 'native_shape' and e.get('shape', {}).get('fill_rgb')]
    # the three pale case cards (01/02/04) must keep their (234,245,254)
    # fill against the (247,250,255) page — outlines stay visible.  A shape
    # qualifies only when its own bbox hugs the card: the old page-sized
    # pale blob merely *covers* every card and averages the contrast away.
    for cx0, cy0, cx1, cy1 in ((137, 499, 984, 1684), (1044, 499, 1886, 1690),
                              (2872, 499, 3709, 1695)):
        hit = next((e for e in fills
                    if max(abs(e['bbox'][0] - cx0), abs(e['bbox'][2] - cx1),
                           abs(e['bbox'][1] - cy0), abs(e['bbox'][3] - cy1)) <= 90
                    and max(abs(a - b) for a, b in zip(
                        e['shape']['fill_rgb'], (234, 245, 254))) <= 8), None)
        assert hit is not None, \
            (cx0, [(e['bbox'], e['shape'].get('fill_rgb')) for e in fills])


def test_bany09_run04_p02_no_fabricated_dark_in_card_gap(tmp_path):
    from image2editable.layout_first import prepare_layout
    source, items = _bany06_fixture('run04', 2)
    slide, _ = prepare_layout(source, tmp_path / 'assets', text_items=items,
                              background_mode='source',
                              label_reader=lambda *a, **k: ([], None))
    # the source strip between the two cards holds only light-blue ink —
    # every dark pixel a PNG carries there is a compositing artefact that
    # renders as a grey dash over the border lines
    for x0, y0, x1, y1 in ((300, 1596, 3680, 1610), (300, 1636, 3680, 1652)):
        for element in slide['visual_elements']:
            c = element.get('component')
            if not c:
                continue
            ix0, iy0 = max(x0, c['x']), max(y0, c['y'])
            ix1, iy1 = min(x1, c['x'] + c['w']), min(y1, c['y'] + c['h'])
            if ix1 <= ix0 or iy1 <= iy0:
                continue
            rgba = np.asarray(Image.open(c['path']).convert('RGBA'))
            zone = rgba[iy0 - c['y']:iy1 - c['y'], ix0 - c['x']:ix1 - c['x']]
            dark = (zone[:, :, :3].max(axis=2) < 160) & (zone[:, :, 3] > 100)
            assert not np.any(dark), \
                (element['object_id'], int(dark.sum()))


def test_bany09_residual_merge_keeps_partial_alpha_colour(tmp_path):
    # run04 p02: the between-cards strip absorbed anti-aliased line-edge
    # fragments whose ~50% alpha was blended onto the empty canvas as
    # colour*alpha — a premultiplied grey that rendered as dashes on the
    # border lines.  A merge must composite properly: the fragment keeps
    # its own colour under union alpha.
    from image2editable.layout_first import _emit_residual
    difference = np.zeros((600, 800), np.uint8)
    text = np.zeros((600, 800), np.uint8)
    host = np.zeros((100, 200, 4), np.uint8)
    host[20:60, 20:180] = (240, 246, 252, 255)         # opaque strip core
    frag = np.zeros((20, 20, 4), np.uint8)
    frag[5:15, 5:15] = (214, 234, 247, 120)            # half-alpha edge sliver
    records = [
        {'rgba': host, 'x': 100, 'y': 100, 'w': 200, 'h': 100, 'owned': False},
        # inside the host's box but over its transparent corner — where the
        # old formula premultiplied the colour onto rgb 0
        {'rgba': frag, 'x': 130, 'y': 170, 'w': 20, 'h': 20, 'owned': False},
    ]
    elements = []
    _emit_residual(records, elements, tmp_path, difference, (600, 800), text, [])
    assert len(elements) == 1
    merged = np.asarray(Image.open(elements[0]['component']['path']))
    px = merged[80, 40]
    assert px[:3].tolist() == pytest.approx([214, 234, 247], abs=2)
    assert px[3] >= 119


def test_bany09_audit_i1_foreign_colour_ink_not_flagged(tmp_path):
    # flatbg p04: the suppressed yellow '接触性出血' unit's first glyph pokes
    # 8 px left of its OCR box, inside the white '思考：面对这位' item's box
    # edge.  That ink is the neighbour element's own — I1 must only fire on
    # raster re-printing the item's *own* glyph colour, otherwise every
    # boundary-straddling foreign glyph reads as a duplicate text.
    from image2editable.layout_first import _audit_invariants
    src = np.full((120, 400, 3), (26, 78, 187), np.uint8)
    src[40:80, 40:140] = 255            # white '…这位' ink
    src[40:80, 136:190] = (255, 209, 32)  # yellow '接' ink protruding left
    png = np.zeros((120, 400, 4), np.uint8)
    png[40:80, 136:190] = (255, 209, 32, 255)
    path = tmp_path / 'unit_0001.png'
    Image.fromarray(png).save(path)
    elements = [{'route': 'raster_component',
                 'component': {'path': str(path), 'x': 0, 'y': 0, 'w': 400, 'h': 120}}]
    items = [{'box': (30, 30, 110, 60), 'text': '…这位', 'color': '#ffffff'},
             {'box': (190, 30, 150, 60), 'text': '…女性', 'color': '#ffffff'}]
    audit = _audit_invariants(elements, items, (120, 400), src)
    assert audit['i1'] == 0
    # the gate still fires when the PNG re-prints ink of the item's own colour
    src2 = src.copy()
    src2[40:80, 136:190] = 255
    png2 = np.zeros((120, 400, 4), np.uint8)
    png2[40:80, 136:190] = (255, 255, 255, 255)
    path2 = tmp_path / 'unit_0002.png'
    Image.fromarray(png2).save(path2)
    elements[0]['component'] = {'path': str(path2), 'x': 0, 'y': 0, 'w': 400, 'h': 120}
    audit = _audit_invariants(elements, items, (120, 400), src2)
    assert audit['i1'] == 1
