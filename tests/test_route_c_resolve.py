"""Route C fallback-page resolution: request binding, donor contract,
draft preservation and idempotency — all offline, no COM needed."""
from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from image2editable import route_c_resolve


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (32, 24), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def run_dir(tmp_path: Path) -> Path:
    """A minimal run root with one bound fallback request."""
    route_dir = tmp_path / "pages" / "page_001" / "route-c"
    route_dir.mkdir(parents=True)
    source = route_dir / "source.png"
    source.write_bytes(_png_bytes())
    quality = route_dir / "quality-report.json"
    quality.write_text('{"rejected": true}', encoding="utf-8")
    request = {
        "schema_version": 1,
        "page_id": "page_001",
        "input_index": 1,
        "target_route": "A",
        "reason": "no_quality_improvement",
        "source_ref": {
            "path": "pages/page_001/route-c/source.png",
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "quality_ref": {
            "path": "pages/page_001/route-c/quality-report.json",
            "sha256": hashlib.sha256(quality.read_bytes()).hexdigest(),
        },
        "repair_round": 5,
        "unresolved_violations": ["duplicate_shadow"],
        "failed_component_ids": ["parent_0008"],
        "status": "awaiting_host",
    }
    (route_dir / "fallback-request.json").write_text(
        json.dumps(request), encoding="utf-8"
    )
    return tmp_path


def _request(run_dir: Path) -> Path:
    return run_dir / "pages" / "page_001" / "route-c" / "fallback-request.json"


def _draft(tmp_path: Path, source_png: Path) -> Path:
    """Hybrid draft: one slide whose only shape is the flattened pic."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        str(source_png), 0, 0,
        width=prs.slide_width, height=prs.slide_height,
    )
    deck = tmp_path / "draft.pptx"
    prs.save(deck)
    return deck


def _donor(tmp_path: Path, *, text: str = "重建文本", color: str = "E8B44C",
           bold: bool = True) -> Path:
    """Accepted single-slide donor: explicit RGB text + a picture."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(100000), Emu(100000),
                                   Emu(2000000), Emu(400000))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(28)
    run.font.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(color)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string("10233F")
    slide.shapes.add_picture(
        str(_png(tmp_path)), Emu(500000), Emu(500000),
        width=Emu(1000000), height=Emu(750000),
    )
    deck = tmp_path / "donor.pptx"
    prs.save(deck)
    return deck


def _png(tmp_path: Path) -> Path:
    png = tmp_path / "donor-obj.png"
    if not png.exists():
        png.write_bytes(_png_bytes())
    return png


def test_resolve_happy_path(run_dir, tmp_path):
    draft = _draft(tmp_path, _png(tmp_path))
    donor = _donor(tmp_path)
    out = tmp_path / "mixed.pptx"
    resolution_path = tmp_path / "resolution.json"
    result = route_c_resolve.resolve_fallback_page(
        _request(run_dir), draft, donor, out,
        resolution_path=resolution_path,
    )
    assert out.is_file() and resolution_path.is_file()
    assert result["status"] == "resolved"
    assert result["slide_part"] == "ppt/slides/slide1.xml"
    doc = json.loads(resolution_path.read_text())
    assert doc["output_ref"]["sha256"] == hashlib.sha256(
        out.read_bytes()
    ).hexdigest()
    prs = Presentation(str(out))
    slide = prs.slides[0]
    texts = [
        "".join(r.text for p in s.text_frame.paragraphs for r in p.runs)
        for s in slide.shapes if s.has_text_frame
    ]
    assert "重建文本" in texts


def test_resolve_rejects_tampered_source(run_dir, tmp_path):
    (run_dir / "pages/page_001/route-c/source.png").write_bytes(b"tampered")
    with pytest.raises(route_c_resolve.HandoffError, match="hash"):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir),
            _draft(tmp_path, _png(tmp_path)),
            _donor(tmp_path),
            tmp_path / "out.pptx",
        )


def test_resolve_rejects_resolved_request(run_dir, tmp_path):
    req = _request(run_dir)
    doc = json.loads(req.read_text())
    doc["status"] = "resolved"
    req.write_text(json.dumps(doc))
    with pytest.raises(route_c_resolve.HandoffError, match="awaiting_host"):
        route_c_resolve.resolve_fallback_page(
            req, _draft(tmp_path, _png(tmp_path)), _donor(tmp_path),
            tmp_path / "out.pptx",
        )


def test_resolve_rejects_donor_with_external_link(run_dir, tmp_path):
    donor = _donor(tmp_path)
    # Inject an external relationship into the donor slide rels.
    fixed = tmp_path / "donor-ext.pptx"
    with zipfile.ZipFile(donor) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    rels_name = "ppt/slides/_rels/slide1.xml.rels"
    rels = etree.fromstring(members[rels_name])
    rel = etree.SubElement(rels, f"{{{route_c_resolve.PR}}}Relationship")
    rel.set("Id", "rIdExt")
    rel.set("Type", "http://example.com/external")
    rel.set("Target", "https://example.com/x.png")
    rel.set("TargetMode", "External")
    members[rels_name] = etree.tostring(rels, xml_declaration=True,
                                        standalone=True)
    with zipfile.ZipFile(fixed, "w") as archive:
        for n, blob in members.items():
            archive.writestr(n, blob)
    with pytest.raises(route_c_resolve.DonorError, match="[Ee]xternal"):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir), _draft(tmp_path, _png(tmp_path)), fixed,
            tmp_path / "out.pptx",
        )


def test_resolve_rejects_theme_font_donor(run_dir, tmp_path):
    donor = _donor(tmp_path)
    fixed = tmp_path / "donor-theme.pptx"
    with zipfile.ZipFile(donor) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    slide = members["ppt/slides/slide1.xml"].replace(
        b'typeface="Arial"', b'typeface="+mn-lt"'
    )
    members["ppt/slides/slide1.xml"] = slide
    with zipfile.ZipFile(fixed, "w") as archive:
        for n, blob in members.items():
            archive.writestr(n, blob)
    with pytest.raises(route_c_resolve.DonorError, match="theme|missing"):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir), _draft(tmp_path, _png(tmp_path)), fixed,
            tmp_path / "out.pptx",
        )


def test_resolve_never_overwrites_existing_output(run_dir, tmp_path):
    out = tmp_path / "mixed.pptx"
    out.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir), _draft(tmp_path, _png(tmp_path)),
            _donor(tmp_path), out,
        )
    assert out.read_bytes() == b"existing"


def test_resolve_out_of_range_index(run_dir, tmp_path):
    req = _request(run_dir)
    doc = json.loads(req.read_text())
    doc["input_index"] = 9
    req.write_text(json.dumps(doc))
    with pytest.raises(route_c_resolve.HandoffError, match="out of range"):
        route_c_resolve.resolve_fallback_page(
            req, _draft(tmp_path, _png(tmp_path)), _donor(tmp_path),
            tmp_path / "out.pptx",
        )


def test_text_contract_enforced(run_dir, tmp_path):
    draft = _draft(tmp_path, _png(tmp_path))
    donor = _donor(tmp_path, text="标题A", color="E8B44C", bold=True)
    contract = tmp_path / "contract.json"
    contract.write_text(json.dumps({
        "items": [{"id": "t1", "text": "标题A", "bold": True,
                   "italic": False, "color": "E8B44C"}]
    }))
    out = tmp_path / "mixed.pptx"
    route_c_resolve.resolve_fallback_page(
        _request(run_dir), draft, donor, out, text_contract_path=contract,
    )
    # Mismatched contract rejected on the same inputs.
    bad = tmp_path / "contract-bad.json"
    bad.write_text(json.dumps({
        "items": [{"id": "t1", "text": "标题A", "bold": False,
                   "italic": False, "color": "E8B44C"}]
    }))
    with pytest.raises(route_c_resolve.DonorError, match="bold"):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir), draft, donor, tmp_path / "o2.pptx",
            text_contract_path=bad,
        )


def _partial_draft(tmp_path: Path, source_png: Path) -> Path:
    """Partial-like draft: full-page background pic + component pic +
    native text — the shape profile resolve must replace wholesale."""
    prs = Presentation()
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(5143500)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.shapes.add_picture(
        str(source_png), 0, 0,
        width=prs.slide_width, height=prs.slide_height,
    )
    background._element.nvPicPr.cNvPr.set("descr", "background.png")
    component = slide.shapes.add_picture(
        str(_png(tmp_path)), Emu(300000), Emu(400000),
        width=Emu(800000), height=Emu(600000),
    )
    component._element.nvPicPr.cNvPr.set("descr", "parent_0002.png")
    box = slide.shapes.add_textbox(
        Emu(100000), Emu(80000), Emu(3000000), Emu(500000)
    )
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "原生标题"
    run.font.size = Pt(28)
    deck = tmp_path / "draft-partial.pptx"
    prs.save(deck)
    return deck


def _slide_xml(pptx_path: Path, slide_part: str = "ppt/slides/slide1.xml"):
    with zipfile.ZipFile(pptx_path) as archive:
        return archive.read(slide_part)


def test_resolve_partial_page_replaces_whole_slide(run_dir, tmp_path):
    draft = _partial_draft(tmp_path, _png(tmp_path))
    donor = _donor(tmp_path)
    out = tmp_path / "resolved.pptx"
    result = route_c_resolve.resolve_fallback_page(
        _request(run_dir), draft, donor, out,
    )
    assert result["resolution_mode"] == "replace"
    assert result["replaced_picture_id"] is None
    slide = Presentation(str(out)).slides[0]
    texts = [
        "".join(r.text for p in s.text_frame.paragraphs for r in p.runs)
        for s in slide.shapes if s.has_text_frame
    ]
    assert "重建文本" in texts
    assert "原生标题" not in texts
    # The component picture and the old background picture are gone;
    # only donor shapes remain.
    pic_descrs = [
        s._element.nvPicPr.cNvPr.get("descr")
        for s in slide.shapes if s.shape_type == 13
    ]
    assert "parent_0002.png" not in pic_descrs
    assert "background.png" not in pic_descrs


def test_resolve_partial_page_preserves_other_slides(run_dir, tmp_path):
    draft = _partial_draft(tmp_path, _png(tmp_path))
    # Append a second, unrelated slide which must survive byte-identical.
    with zipfile.ZipFile(draft) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    prs = Presentation(str(draft))
    extra = prs.slides.add_slide(prs.slide_layouts[6])
    extra.shapes.add_textbox(0, 0, Emu(1000000), Emu(500000)).text = (
        "keep me"
    )
    prs.save(draft)
    donor = _donor(tmp_path)
    out = tmp_path / "resolved.pptx"
    route_c_resolve.resolve_fallback_page(
        _request(run_dir), draft, donor, out,
    )
    slide2 = Presentation(str(out)).slides[1]
    assert slide2.shapes[0].text_frame.text == "keep me"


def _donor_with_gradient(tmp_path: Path, *, scheme: bool) -> Path:
    donor = _donor(tmp_path)
    fixed = tmp_path / ("donor-grad-scheme.pptx" if scheme
                      else "donor-grad.pptx")
    color0 = "schemeClr" if scheme else "srgbClr"
    sp = (
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/'
        'presentationml/2006/main" xmlns:a="http://schemas.openxmlformats'
        '.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="90" name="gr"/>'
        '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>'
        '<a:off x="0" y="0"/><a:ext cx="500000" cy="500000"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:gradFill>'
        f'<a:gsLst><a:gs pos="0"><a:{color0} val="10233F"/></a:gs>'
        '<a:gs pos="100000"><a:srgbClr val="E8B44C"/></a:gs></a:gsLst>'
        '<a:lin ang="5400000" scaled="1"/></a:gradFill></p:spPr>'
        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
    )
    with zipfile.ZipFile(donor) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    slide = members["ppt/slides/slide1.xml"].decode("utf-8")
    slide = slide.replace("</p:spTree>", sp + "</p:spTree>")
    members["ppt/slides/slide1.xml"] = slide.encode("utf-8")
    with zipfile.ZipFile(fixed, "w") as archive:
        for name, blob in members.items():
            archive.writestr(name, blob)
    return fixed


def test_donor_gradient_with_explicit_stops_accepted(run_dir, tmp_path):
    donor = _donor_with_gradient(tmp_path, scheme=False)
    out = tmp_path / "mixed.pptx"
    result = route_c_resolve.resolve_fallback_page(
        _request(run_dir), _draft(tmp_path, _png(tmp_path)), donor, out,
    )
    assert result["status"] == "resolved"
    assert out.is_file()


def test_donor_gradient_with_scheme_stop_rejected(run_dir, tmp_path):
    donor = _donor_with_gradient(tmp_path, scheme=True)
    with pytest.raises(
        route_c_resolve.DonorError, match="non-explicit"
    ):
        route_c_resolve.resolve_fallback_page(
            _request(run_dir), _draft(tmp_path, _png(tmp_path)), donor,
            tmp_path / "out.pptx",
        )


def test_partial_background_swap(run_dir, tmp_path):
    draft = _partial_draft(tmp_path, _png(tmp_path))
    new_bg = tmp_path / "regenerated-bg.png"
    buf = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (64, 36), (10, 200, 90)).save(buf, format="PNG")
    new_bg.write_bytes(buf.getvalue())
    out = tmp_path / "bg-swapped.pptx"
    result = route_c_resolve.resolve_partial_background(
        _request(run_dir), draft, new_bg, out,
        resolution_path=tmp_path / "resolution.json",
    )
    assert result["status"] == "resolved"
    assert result["resolution_kind"] == "background"
    # Slide XML untouched; only rels repointed + new media part added.
    assert _slide_xml(out) == _slide_xml(draft)
    with zipfile.ZipFile(out) as archive:
        assert archive.read("ppt/media/bg_resolve_1.png") == (
            new_bg.read_bytes()
        )
        rels = etree.fromstring(
            archive.read("ppt/slides/_rels/slide1.xml.rels")
        )
        targets = {
            r.get("Id"): r.get("Target")
            for r in rels.findall(
                f"{{{route_c_resolve.PR}}}Relationship"
            )
        }
        assert "bg_resolve_1.png" in targets["rId2"]
    slide = Presentation(str(out)).slides[0]
    assert len(slide.shapes) == 3


def test_partial_background_rejects_flattened_page(run_dir, tmp_path):
    new_bg = tmp_path / "bg.png"
    new_bg.write_bytes(_png_bytes())
    with pytest.raises(
        route_c_resolve.HandoffError, match="partial"
    ):
        route_c_resolve.resolve_partial_background(
            _request(run_dir), _draft(tmp_path, _png(tmp_path)),
            new_bg, tmp_path / "out.pptx",
        )


def test_partial_background_rejects_non_image(run_dir, tmp_path):
    bogus = tmp_path / "bg.bin"
    bogus.write_bytes(b"not an image at all")
    with pytest.raises(route_c_resolve.HandoffError, match="PNG or JPEG"):
        route_c_resolve.resolve_partial_background(
            _request(run_dir), _partial_draft(tmp_path, _png(tmp_path)),
            bogus, tmp_path / "out.pptx",
        )
